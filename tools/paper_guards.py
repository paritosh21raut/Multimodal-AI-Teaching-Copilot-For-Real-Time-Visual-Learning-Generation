"""Paper: what the coverage guard contributes, and how often the copilot rewrites shown text, on D2 (zero LLM tokens).

    .venv/Scripts/python tools/paper_guards.py        # -> verify_results/_eval/guards_d2.json

Coverage guard. Every live activation is logged ("model left out N spoken sentence(s); shown as said: [...]").
Each added sentence is matched to its transcript line (word overlap); a key fact the annotators marked Y/P on such a
line reached the slide through the guard, so coverage without the guard counts it as missed (the guard fires only
when the LLM output shares <= 1/3 of the sentence's words, so the fact was otherwise absent). Same aggregation as
`eval_sheets.py score` (mean over lectures x annotators).

Churn. Text of a shown slide version that is missing from the next version of the same slide (the metric of
`baseline_llm_slides.py --metrics`), over every annotated lecture, with the cause of each loss: a teacher command,
an applied correction (concern), or neither.
"""
from __future__ import annotations

import ast
import json
import re
import sqlite3
import sys
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parent))
from baseline_llm_slides import _texts  # noqa: E402
from eval_sheets import ANNOTATE, ANNOTATORS, EVAL, read_workbook  # noqa: E402
from eval_test import OUT  # noqa: E402

WORD = re.compile(r"[a-z0-9]+")
RNG = np.random.default_rng(2026)


def words(s: str) -> set[str]:
    return {w[:6] for w in WORD.findall(s.lower()) if len(w) > 2}


def guard_sentences(folder: Path) -> list[str]:
    out = []
    for line in (folder / "terminal_log.txt").read_text(encoding="utf-8", errors="replace").splitlines():
        m = re.search(r"model left out \d+ spoken sentence\(s\); shown as said: (\[.*\])", line)
        if m:
            out += ast.literal_eval(m.group(1))
    return out


def transcript(folder: Path) -> list[str]:
    from openpyxl import load_workbook
    wb = load_workbook(EVAL / "annotation" / "A" / f"{folder.name}_A.xlsx", data_only=True)
    return [str(r[2] or "") for r in list(wb["Transcript"].iter_rows(min_row=2, values_only=True))]


def guard_lines(folder: Path) -> tuple[set[int], int]:
    lines = [words(t) for t in transcript(folder)]
    hit = set()
    sents = guard_sentences(folder)
    for s in sents:
        w = words(s)
        best = max(range(len(lines)), key=lambda i: len(w & lines[i]) / max(1, len(w)))
        if len(w & lines[best]) / max(1, len(w)) >= 0.6:
            hit.add(best)  # 0-based row of the Transcript sheet
    return hit, len(sents)


def churn(folder: Path) -> dict:
    db = sqlite3.connect(folder / "session" / "session.sqlite")
    ev = [(t, ty, json.loads(p)) for t, ty, p in db.execute("select ts, type, payload from events order by seq")]
    last: dict[str, list[str]] = {}
    shown = lost = by_teacher = by_correction = 0
    since_cmd = since_fix = False
    for _, ty, p in ev:
        if ty == "CommandReceived":
            since_cmd = True
        elif ty in ("ConcernApplied", "CorrectionApplied") or (ty == "ConcernRaised" and p.get("applied")):
            since_fix = True
        elif ty == "SlidePatch" and p["spec"].get("layout") != "title":
            cur = _texts(p["spec"])
            prev = last.get(p["slide_id"])
            if prev is not None:
                gone = [x for x in prev if x not in cur]
                shown += len(prev)
                lost += len(gone)
                if gone and since_cmd:
                    by_teacher += len(gone)
                elif gone and since_fix:
                    by_correction += len(gone)
            last[p["slide_id"]] = cur
            since_cmd = since_fix = False
    return dict(shown=shown, lost=lost, by_teacher=by_teacher, by_correction=by_correction)


def main() -> None:
    rows, ch = [], []
    for name in ANNOTATE:
        folder = OUT / name
        hit, n_sent = guard_lines(folder)
        c = churn(folder)
        ch.append(c)
        for w in ANNOTATORS:
            a = read_workbook(EVAL / "annotation" / w / f"{name}_{w}.xlsx")
            from openpyxl import load_workbook
            tr = list(load_workbook(EVAL / "annotation" / w / f"{name}_{w}.xlsx", data_only=True)["Transcript"]
                      .iter_rows(min_row=2, values_only=True))
            facts = [(i, str(r[5]).strip().upper()[:1] if r[5] not in (None, "") else None)
                     for i, r in enumerate(tr) if r[4] not in (None, "")]
            facts = [(i, x) for i, x in facts if x]
            if not facts or a is None:
                continue
            val = {"Y": 1.0, "P": 0.5}
            with_g = sum(val.get(x, 0) for _, x in facts) / len(facts)
            no_g = sum(val.get(x, 0) for i, x in facts if i not in hit) / len(facts)
            via = sum(1 for i, x in facts if i in hit and x in val)
            rows.append(dict(lecture=name, annotator=w, facts=len(facts), via_guard=via, coverage=with_g,
                             coverage_without_guard=no_g, guard_sentences=n_sent))
    cov = np.array([r["coverage"] for r in rows])
    cov0 = np.array([r["coverage_without_guard"] for r in rows])
    diff = cov - cov0
    boots = [diff[RNG.integers(0, len(diff), len(diff))].mean() for _ in range(2000)]
    tot = {k: sum(c[k] for c in ch) for k in ch[0]}
    res = dict(
        coverage=round(float(cov.mean()), 4), coverage_without_guard=round(float(cov0.mean()), 4),
        gain_pp=round(100 * float(diff.mean()), 2),
        gain_ci_pp=[round(100 * float(np.percentile(boots, q)), 2) for q in (2.5, 97.5)],
        facts=sum(r["facts"] for r in rows), facts_via_guard=sum(r["via_guard"] for r in rows),
        guard_sentences=sum(r["guard_sentences"] for r in rows if r["annotator"] == "A"),
        churn=dict(tot, lost_share=round(tot["lost"] / max(1, tot["shown"]), 4)),
        per_lecture=rows, churn_per_lecture=dict(zip(ANNOTATE, ch)))
    (EVAL / "guards_d2.json").write_text(json.dumps(res, indent=1), encoding="utf-8")
    print({k: v for k, v in res.items() if k not in ("per_lecture", "churn_per_lecture")})


if __name__ == "__main__":
    main()
