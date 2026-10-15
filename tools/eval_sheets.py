"""Paper validation, phases 3 + 4 (zero LLM tokens): annotation workbooks, teacher rating form, scoring.

    .venv/Scripts/python tools/eval_sheets.py sheets     # -> verify_results/_eval/ (annotation/A, annotation/B, rating/)
    .venv/Scripts/python tools/eval_sheets.py score      # filled sheets + returned ratings -> verify_results/_eval/results.json

Lectures are the folders written by tools/eval_test.py (verify_results/testN, webN). Plan: docs/paper/VALIDATION.md.

Units of annotation are the recognised utterances (TranscriptFinal, numbered 1..n). The system's topic boundaries are
the utterances whose interpretation unit opened a new content slide that is still in the final deck. Overflow parts
(part > 1 with the same facet as the slide before) are not boundaries; the end-of-lecture summary / key-concept slides
(origin "materials") are not lecture slides.
"""
from __future__ import annotations

import base64
import csv
import io
import json
import random
import re
import sqlite3
import statistics as st
import sys
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parent))
from eval_test import OUT, block_text, deck_from_log  # noqa: E402

EVAL = OUT / "_eval"
SEED = 2026
ANNOTATORS = ("A", "B")
# 18 lectures for annotation: 12 recorded + 6 public, mixed subjects, 6-17 min (test8/27/30 too long, <5 min too short)
ANNOTATE = ["test3", "test10", "test14", "test15", "test17", "test20",       # recorded
            "test21", "test22", "test23", "test24", "test25", "test29",      # recorded
            "web1", "web2", "web3", "web4", "web5", "web6"]                  # public recorded lectures
RATE_PER_GROUP = 10
CRITERIA = [("correctness", "The content on the slide is factually correct."),
            ("relevance", "The slide shows what the teacher said in the given part of the lecture."),
            ("readability", "The slide is easy to read from the back of a classroom (amount of text, layout)."),
            ("usefulness", "The slide would help students follow and learn this part of the lecture.")]


# ---------------------------------------------------------------------------------------------- lecture data
def lectures() -> list[Path]:
    key = lambda p: (p.name.rstrip("0123456789"), int(re.sub(r"\D", "", p.name) or 0))  # noqa: E731
    return sorted((p for p in OUT.iterdir() if p.is_dir() and (p / "meta.json").exists()
                   and (p / "session" / "session.sqlite").exists()), key=key)


def group(meta: dict) -> str:
    # The 30 recorded tests were given by six speakers; the per-test speaker labels are not reliable, so all of them
    # form one group. The web lectures are the public held-out group.
    return "public" if meta.get("source") == "web" else "recorded"


def group_of(label: str) -> str:
    """Map a group label stored earlier (rating_key.json: speaker A/B, web) to the merged groups."""
    return "public" if label in ("web", "public") else "recorded"


class Lecture:
    def __init__(self, folder: Path):
        self.name, self.folder = folder.name, folder
        self.meta = json.loads((folder / "meta.json").read_text(encoding="utf-8"))
        self.group = group(self.meta)
        db = folder / "session" / "session.sqlite"
        ev = [(t, json.loads(p)) for t, p in sqlite3.connect(db).execute(
            "select type, payload from events order by seq")]
        self.utts = [p["segment"] for t, p in ev if t == "TranscriptFinal"]
        self.idx = {u["id"]: i for i, u in enumerate(self.utts, 1)}
        self.deck = [s for s in deck_from_log(db)]
        self.content = [(n, s) for n, s in enumerate(self.deck, 1)
                        if s.get("origin", "lecture") == "lecture" and s.get("layout") != "title"]
        # which utterances fed each slide: patches are attributed to the interpretation that came just before them
        self.feed: dict[str, list[int]] = {}
        self.opened: dict[str, int] = {}
        segs: list[str] = []
        for t, p in ev:
            if t == "InterpretationReady":
                segs = p.get("segment_ids", [])
            elif t == "SlidePatch" and segs:
                lines = sorted(self.idx[s] for s in segs if s in self.idx)
                if not lines:
                    continue
                f = self.feed.setdefault(p["slide_id"], [])
                f += [i for i in lines if i not in f]
                if p.get("op") == "add":
                    self.opened.setdefault(p["slide_id"], lines[0])
        self.ir = [p for t, p in ev if t == "InterpretationReady"]
        self.prompt_tokens = sum(p.get("prompt_tokens", 0) for t, p in ev if t == "InterpretRequested")
        self.commands = [p.get("command", {}).get("kind") for t, p in ev if t == "CommandReceived"]

    def boundaries(self) -> list[int]:
        """Utterance numbers where the system opened a new content slide (first slide included)."""
        out, prev = set(), None
        for _, s in self.content:
            overflow = (s.get("part") or 1) > 1 and prev is not None and s.get("facet") == prev.get("facet")
            if not overflow and s["id"] in self.opened:
                out.add(self.opened[s["id"]])
            prev = s
        return sorted(out)

    def statements(self) -> list[tuple[int, str, str]]:
        rows = []
        for n, s in self.content:
            for b in s.get("blocks", []):
                if b.get("type") == "image":
                    continue
                for line in block_text(b):
                    text = re.sub(r"[*_]", "", line).strip().lstrip("- ").strip()
                    if text and not re.fullmatch(r"\[[a-z_, ]+\]", text):
                        rows.append((n, s.get("title", ""), text))
        return rows


# ---------------------------------------------------------------------------------------------- annotation workbooks
INSTRUCTIONS = """Annotation of lecture {name}  (annotator {who})

Work ALONE: do not look at the other annotator's sheet and do not discuss until both are finished.
Do steps 1 and 2 BEFORE you look at the slides.

Sheet "Transcript" (one row per recognised utterance, as the system heard it):
 1. New topic? — put 1 in the row where a NEW TOPIC OR SUBTOPIC starts, i.e. where a teacher writing on a board would
    start a new board / slide. Mark the first content line of the lecture too. Leave the column empty otherwise.
    Small digressions, jokes and organisational talk are not new topics.
 2. Key fact — in the row where it is said, write each key fact a student should take away (a short phrase, e.g.
    "ionic bond = complete electron transfer"). About 1 key fact per minute of lecture; repeated facts only once.
 3. Now open the slides: verify_results/{name}/slides/NN.png (or slides.md). For every key fact you wrote, fill
    On slides?:  Y = shown on a slide,  P = partly / vaguely,  N = not shown.

Sheet "Slide statements" (every statement the system put on a lecture slide; images left out):
 4. Supported?:  Y = the teacher said this (or it follows directly from what was said),  P = partly,
    N = not said by the teacher (added by the system).  Use the Transcript sheet / transcript.txt to check.
 5. Wrong?: put 1 if the statement is factually wrong (regardless of whether the teacher said it).
 Comment is optional.

Speech recognition errors in the transcript are part of the system; judge the meaning the teacher intended.
Save the file with the same name when finished.
"""


def write_workbook(lec: Lecture, who: str, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    for r, line in enumerate(INSTRUCTIONS.format(name=lec.name, who=who).splitlines(), 1):
        ws.cell(r, 1, line)
    ws.column_dimensions["A"].width = 120
    head = Font(bold=True)
    fill = PatternFill("solid", fgColor="FFF2CC")

    def sheet(title, cols, widths, rows, fill_cols):
        sh = wb.create_sheet(title)
        sh.append(cols)
        for c in sh[1]:
            c.font = head
        for row in rows:
            sh.append(list(row))
        for i, w in enumerate(widths):
            sh.column_dimensions[chr(65 + i)].width = w
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if c.column - 1 in fill_cols:
                    c.fill = fill
        sh.freeze_panes = "A2"

    sheet("Transcript", ["#", "time (s)", "what was said", "New topic? (1)", "Key fact said here", "On slides? (Y/P/N)"],
          [6, 9, 90, 14, 40, 16],
          [(i, round(u["start"], 1), " ".join(u["text"].split()), None, None, None) for i, u in enumerate(lec.utts, 1)],
          {3, 4, 5})
    sheet("Slide statements", ["#", "slide", "slide title", "statement", "Supported? (Y/P/N)", "Wrong? (1)", "comment"],
          [6, 7, 30, 80, 18, 11, 30],
          [(i, n, title, text, None, None, None) for i, (n, title, text) in enumerate(lec.statements(), 1)],
          {4, 5, 6})
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------------------------------------------------------------------------------------------- rating form
def sample_slides(lecs: list[Lecture]) -> list[dict]:
    rng = random.Random(SEED)
    items = []
    for g in sorted({lec.group for lec in lecs}):
        pool = [lec for lec in lecs if lec.group == g]
        rng.shuffle(pool)
        picks: list[tuple[Lecture, int]] = []
        per = 1
        while len(picks) < RATE_PER_GROUP and per <= 5:
            for lec in pool:
                cands = [(n, s) for n, s in lec.content if lec.feed.get(s["id"]) and (lec, n) not in picks]
                if cands and sum(1 for x, _ in picks if x is lec) < per and len(picks) < RATE_PER_GROUP:
                    picks.append((lec, rng.choice(cands)[0]))
            per += 1
        for lec, n in picks:
            s = lec.deck[n - 1]
            lines = lec.feed[s["id"]]
            items.append(dict(group=g, lecture=lec.name, slide=n, slide_id=s["id"], title=s.get("title", ""),
                              lines=[lines[0], lines[-1]],
                              span=[round(lec.utts[lines[0] - 1]["start"]), round(lec.utts[lines[-1] - 1]["end"])],
                              said=" ".join(" ".join(lec.utts[i - 1]["text"].split()) for i in lines)))
    rng.shuffle(items)
    for k, it in enumerate(items, 1):
        it["item"] = k
    return items


def jpeg_b64(png: Path) -> str:
    from PIL import Image

    im = Image.open(png).convert("RGB")
    im.thumbnail((1280, 720))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=82)
    return base64.b64encode(buf.getvalue()).decode()


def rating_html(items: list[dict]) -> str:
    crit_head = "".join(f"<li><b>{c.capitalize()}</b>: {d}</li>" for c, d in CRITERIA)
    blocks = []
    for it in items:
        said = it["said"] if len(it["said"]) <= 1400 else it["said"][:1400] + " …"
        m0, m1 = divmod(it["span"][0], 60), divmod(it["span"][1], 60)
        radios = "".join(
            f"<tr><td>{c.capitalize()}</td>" + "".join(
                f'<td><label><input type="radio" name="i{it["item"]}_{c}" value="{v}"> {v}</label></td>'
                for v in range(1, 6)) + "</tr>" for c, _ in CRITERIA)
        blocks.append(f"""
<section class="item" id="item{it['item']}">
  <h2>Slide {it['item']} of {len(items)}</h2>
  <img alt="slide {it['item']}" src="data:image/jpeg;base64,{jpeg_b64(OUT / it['lecture'] / 'slides' / f"{it['slide']:02d}.png")}">
  <details open><summary>What the teacher said for this slide ({m0[0]}:{m0[1]:02d}–{m1[0]}:{m1[1]:02d} of the lecture,
  as recognised)</summary><p class="said">{said}</p></details>
  <table class="r"><tr><th></th><th>1 poor</th><th>2</th><th>3</th><th>4</th><th>5 very good</th></tr>{radios}</table>
  <textarea name="i{it['item']}_comment" placeholder="Comment (optional)"></textarea>
</section>""")
    names = json.dumps([f"i{it['item']}_{c}" for it in items for c, _ in CRITERIA] +
                       [f"i{it['item']}_comment" for it in items])
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1"><title>Slide Rating Form</title>
<style>
:root{{--bg:#fafaf7;--fg:#1d1d1b;--mut:#5c5c57;--card:#fff;--line:#d9d9d2;--acc:#1f5fa8}}
@media (prefers-color-scheme: dark){{:root{{--bg:#161615;--fg:#ecebe6;--mut:#a3a29b;--card:#20201e;--line:#3a3a36;--acc:#7fb0ea}}}}
body{{background:var(--bg);color:var(--fg);font:16px/1.5 Georgia,serif;margin:0 auto;max-width:980px;padding:16px}}
h1{{font-size:1.6em}} .item{{background:var(--card);border:1px solid var(--line);border-radius:8px;padding:16px;margin:20px 0}}
img{{width:100%;height:auto;border:1px solid var(--line)}} .said{{color:var(--mut);font-size:.92em}}
table.r{{border-collapse:collapse;margin:10px 0;width:100%}} table.r td,table.r th{{padding:4px 6px;text-align:center;border-bottom:1px solid var(--line)}}
table.r td:first-child{{text-align:left;font-weight:bold}} textarea{{width:100%;min-height:48px;box-sizing:border-box}}
.bar{{position:sticky;top:0;background:var(--bg);padding:8px 0;border-bottom:1px solid var(--line)}}
button{{font:inherit;padding:6px 14px;background:var(--acc);color:#fff;border:0;border-radius:6px;cursor:pointer}}
</style></head><body>
<h1>Rating of automatically generated lecture slides</h1>
<p>The slides below were produced live, while a teacher was speaking, by a system that listens to the lecture.
Under each slide you see what the teacher said for it (as the speech recogniser heard it; recognition errors are
possible). Please rate every slide from 1 (poor) to 5 (very good):</p><ul>{crit_head}</ul>
<p>Rate each slide on its own; there are no right answers. Your name is not recorded; choose your rater code
below. Your answers are kept in this browser if you close the page. When finished, press <b>Download my ratings</b>
and send the CSV file back (if another teacher uses this computer after you, press <b>Clear all answers</b>).</p>
<div class="bar">Rater code: <select id="rater"><option>T1</option><option>T2</option><option>T3</option>
<option>T4</option><option>T5</option></select>
&nbsp; <span id="done"></span> &nbsp; <button id="dl">Download my ratings</button>
<button id="clr" title="For the next teacher on this computer">Clear all answers</button></div>
{''.join(blocks)}
<script>
const NAMES = {names};
const KEY = "slide-rating-v1";
function load() {{ try {{ return JSON.parse(localStorage.getItem(KEY) || "{{}}"); }} catch (e) {{ return {{}}; }} }}
function save(s) {{ try {{ localStorage.setItem(KEY, JSON.stringify(s)); }} catch (e) {{}} }}
const state = load();
function count() {{
  const n = NAMES.filter(k => !k.endsWith("_comment") && state[k]).length;
  const all = NAMES.filter(k => !k.endsWith("_comment")).length;
  document.getElementById("done").textContent = n + " of " + all + " ratings given";
}}
for (const [k, v] of Object.entries(state)) {{
  if (k === "rater") {{ document.getElementById("rater").value = v; continue; }}
  if (k.endsWith("_comment")) {{ const t = document.querySelector(`textarea[name="${{k}}"]`); if (t) t.value = v; continue; }}
  const el = [...document.querySelectorAll(`input[name="${{k}}"]`)].find(e => e.value === String(v));
  if (el) el.checked = true;
}}
document.addEventListener("input", e => {{
  const t = e.target;
  if (t.id === "rater") state.rater = t.value; else if (t.name) state[t.name] = t.value;
  save(state); count();
}});
document.getElementById("dl").onclick = () => {{
  const rater = document.getElementById("rater").value;
  const rows = [["rater", "item", "criterion", "score", "comment"]];
  const items = [...new Set(NAMES.map(k => k.split("_")[0]))];
  for (const it of items) for (const c of {json.dumps([c for c, _ in CRITERIA])})
    rows.push([rater, it.slice(1), c, state[it + "_" + c] || "", ""]);
  for (const it of items) if (state[it + "_comment"]) rows.push([rater, it.slice(1), "comment", "", state[it + "_comment"]]);
  const csv = rows.map(r => r.map(x => '"' + String(x).replace(/"/g, '""') + '"').join(",")).join("\\n");
  const a = document.createElement("a");
  a.href = URL.createObjectURL(new Blob(["\\ufeff" + csv], {{type: "text/csv"}}));
  a.download = "ratings_" + rater + ".csv"; a.click();
}};
document.getElementById("clr").onclick = () => {{
  if (!confirm("Clear all answers on this computer? Download your ratings first.")) return;
  try {{ localStorage.removeItem(KEY); }} catch (e) {{}}
  location.reload();
}};
count();
</script></body></html>"""


def make_sheets() -> None:
    lecs = [Lecture(p) for p in lectures()]
    by = {lec.name: lec for lec in lecs}
    missing = [n for n in ANNOTATE if n not in by]
    if missing:
        print("[SHEETS] not recorded yet, skipped:", ", ".join(missing))
    for who in ANNOTATORS:
        for n in ANNOTATE:
            if n in by:
                path = EVAL / "annotation" / who / f"{n}_{who}.xlsx"
                if path.exists():
                    print(f"[SHEETS] kept {path.relative_to(OUT.parent)} (exists; delete it to regenerate)")
                    continue
                write_workbook(by[n], who, path)
    items = sample_slides(lecs)
    rate = EVAL / "rating"
    (rate / "returned").mkdir(parents=True, exist_ok=True)
    (rate / "rating_key.json").write_text(json.dumps(items, indent=1, ensure_ascii=False), encoding="utf-8")
    (rate / "rating_form.html").write_text(rating_html(items), encoding="utf-8")
    groups = {g: sum(1 for it in items if it["group"] == g) for g in sorted({it["group"] for it in items})}
    print(f"[SHEETS] {sum(1 for n in ANNOTATE if n in by)} lectures x {len(ANNOTATORS)} annotators -> "
          f"{(EVAL / 'annotation').relative_to(OUT.parent)}")
    print(f"[SHEETS] rating form: {len(items)} slides {groups} from {len({it['lecture'] for it in items})} lectures -> "
          f"{(rate / 'rating_form.html').relative_to(OUT.parent)}  (returned CSVs go in rating/returned/)")


# ---------------------------------------------------------------------------------------------- scoring
def boundary_prf(ref: list[int], hyp: list[int], tol: int = 1) -> tuple[float, float, float]:
    used, tp = set(), 0
    for h in hyp:
        m = next((r for r in sorted(ref, key=lambda r: abs(r - h)) if abs(r - h) <= tol and r not in used), None)
        if m is not None:
            used.add(m)
            tp += 1
    p = tp / len(hyp) if hyp else 0.0
    r = tp / len(ref) if ref else 0.0
    return p, r, (2 * p * r / (p + r) if p + r else 0.0)


def pk(ref: list[int], hyp: list[int], n: int) -> float:
    """Beeferman P_k over n utterances; boundaries are the first utterance of each segment."""
    def seg_ids(bs):
        bs = set(bs) | {1}
        out, cur = [], 0
        for i in range(1, n + 1):
            cur += i in bs and i != 1
            out.append(cur)
        return out
    r, h = seg_ids(ref), seg_ids(hyp)
    k = max(2, round(n / max(1, len(set(ref) | {1})) / 2))
    if n <= k:
        return float("nan")
    return float(np.mean([(r[i] == r[i + k]) != (h[i] == h[i + k]) for i in range(n - k)]))


def agreement(a: list, b: list) -> float:
    pairs = [(x, y) for x, y in zip(a, b) if x is not None and y is not None]
    return sum(x == y for x, y in pairs) / len(pairs) if pairs else float("nan")


def kappa(a: list, b: list) -> float:
    pairs = [(x, y) for x, y in zip(a, b) if x is not None and y is not None]
    if not pairs:
        return float("nan")
    cats = sorted({c for p in pairs for c in p})
    po = sum(x == y for x, y in pairs) / len(pairs)
    pe = sum((sum(x == c for x, _ in pairs) / len(pairs)) * (sum(y == c for _, y in pairs) / len(pairs)) for c in cats)
    return (po - pe) / (1 - pe) if pe < 1 else 1.0


def read_workbook(path: Path) -> dict | None:
    from openpyxl import load_workbook

    if not path.exists():
        return None
    wb = load_workbook(path, data_only=True)
    norm = lambda v: str(v).strip().upper()[:1] if v not in (None, "") else None  # noqa: E731
    tr = list(wb["Transcript"].iter_rows(min_row=2, values_only=True))
    ss = list(wb["Slide statements"].iter_rows(min_row=2, values_only=True))
    bounds = [r[0] for r in tr if r[3] not in (None, "", 0, "0")]
    facts = [norm(r[5]) for r in tr if r[4] not in (None, "")]
    supported = [norm(r[4]) for r in ss]
    wrong = [r[5] not in (None, "", 0, "0") for r in ss]
    if not bounds and not facts and not any(supported):
        return None  # not filled in yet
    return dict(bounds=bounds, facts=facts, supported=supported, wrong=wrong)


def icc2(m: np.ndarray) -> float:
    """ICC(2,1), two-way random, absolute agreement; m = items x raters, complete rows only."""
    m = m[~np.isnan(m).any(axis=1)]
    n, k = m.shape
    if n < 2 or k < 2:
        return float("nan")
    g = m.mean()
    msr = k * ((m.mean(axis=1) - g) ** 2).sum() / (n - 1)
    msc = n * ((m.mean(axis=0) - g) ** 2).sum() / (k - 1)
    mse = ((m - m.mean(axis=1, keepdims=True) - m.mean(axis=0, keepdims=True) + g) ** 2).sum() / ((n - 1) * (k - 1))
    return float((msr - mse) / (msr + (k - 1) * mse + k * (msc - mse) / n))


def score() -> None:
    lecs = {p.name: Lecture(p) for p in lectures()}
    res: dict = {"dataset": {}, "annotation": {}, "ratings": {}}
    # --- dataset (everything recorded)
    for g in sorted({lec.group for lec in lecs.values()}) + ["all"]:
        sel = [lec for lec in lecs.values() if g in ("all", lec.group)]
        units = [p for lec in sel for p in lec.ir]
        res["dataset"][g] = dict(
            lectures=len(sel), minutes=round(sum(lec.meta["minutes"] for lec in sel), 1),
            utterances=sum(len(lec.utts) for lec in sel), units=len(units),
            fallback_units=sum(1 for p in units if p.get("fallback")),
            slides=sum(len(lec.content) for lec in sel), prompt_tokens=sum(lec.prompt_tokens for lec in sel),
            teacher_deletions=sum(c in ("delete_item", "delete_slide") for lec in sel for c in lec.commands))
    # --- annotation
    per, rows = {}, []
    for n in ANNOTATE:
        if n not in lecs:
            continue
        lec = lecs[n]
        ann = {who: read_workbook(EVAL / "annotation" / who / f"{n}_{who}.xlsx") for who in ANNOTATORS}
        hyp = lec.boundaries()
        r = dict(lecture=n, group=lec.group, utterances=len(lec.utts), system_boundaries=len(hyp),
                 fallback_share=sum(1 for p in lec.ir if p.get("fallback")) / max(1, len(lec.ir)))
        for who, a in ann.items():
            if a is None:
                continue
            p, rc, f = boundary_prf(a["bounds"], hyp)
            facts = [x for x in a["facts"] if x]
            sup = [x for x in a["supported"] if x]
            r[who] = dict(ref_boundaries=len(a["bounds"]), P=p, R=rc, F1=f, Pk=pk(a["bounds"], hyp, len(lec.utts)),
                          facts=len(a["facts"]),
                          coverage=(sum(x == "Y" for x in facts) + 0.5 * sum(x == "P" for x in facts)) / len(facts)
                          if facts else None,
                          statements=len(sup), unsupported=sum(x == "N" for x in sup) / len(sup) if sup else None,
                          wrong=sum(a["wrong"]) / len(a["wrong"]) if a["wrong"] else None,
                          wrong_n=sum(a["wrong"]),
                          wrong_said=sum(1 for w, x in zip(a["wrong"], a["supported"]) if w and x == "Y"))
        if ann["A"] and ann["B"]:
            ba = [int(i in ann["A"]["bounds"]) for i in range(1, len(lec.utts) + 1)]
            bb = [int(i in ann["B"]["bounds"]) for i in range(1, len(lec.utts) + 1)]
            r["kappa_boundary"] = kappa(ba, bb)
            r["kappa_supported"] = kappa(ann["A"]["supported"], ann["B"]["supported"])
            r["_a"], r["_b"] = (ba, ann["A"]["supported"]), (bb, ann["B"]["supported"])
        rows.append(r)
    if rows:
        def mean(key, sel):
            v = [x[w][key] for x in sel for w in ANNOTATORS if w in x and x[w].get(key) is not None
                 and not (isinstance(x[w][key], float) and np.isnan(x[w][key]))]
            return round(float(np.mean(v)), 3) if v else None
        # LLM and fallback reported separately: lectures whose units were mostly interpreted by the LLM vs. mostly by
        # the deterministic fallback (quota spent / provider failure)
        pick = {"llm_lectures": lambda x: x["fallback_share"] < 0.5, "fallback_lectures": lambda x: x["fallback_share"] >= 0.5}
        for g in sorted({x["group"] for x in rows}) + ["all", "llm_lectures", "fallback_lectures"]:
            sel = [x for x in rows if (pick[g](x) if g in pick else g in ("all", x["group"]))]
            if not sel:
                continue
            both = [x for x in sel if "_a" in x]
            per[g] = dict(lectures=len(sel), F1=mean("F1", sel), P=mean("P", sel), R=mean("R", sel),
                          Pk=mean("Pk", sel), coverage=mean("coverage", sel), unsupported=mean("unsupported", sel),
                          wrong=mean("wrong", sel), facts=sum(x[w]["facts"] for x in sel for w in ANNOTATORS if w in x),
                          statements=sum(x[w]["statements"] for x in sel for w in ANNOTATORS if w in x),
                          wrong_n=sum(x[w]["wrong_n"] for x in sel for w in ANNOTATORS if w in x),
                          wrong_said=sum(x[w]["wrong_said"] for x in sel for w in ANNOTATORS if w in x),
                          lecture_names=[x["lecture"] for x in sel],
                          kappa_boundary=round(kappa(sum((x["_a"][0] for x in both), []),
                                                     sum((x["_b"][0] for x in both), [])), 3) if both else None,
                          kappa_supported=round(kappa(sum((x["_a"][1] for x in both), []),
                                                      sum((x["_b"][1] for x in both), [])), 3) if both else None,
                          agreement_boundary=round(agreement(sum((x["_a"][0] for x in both), []),
                                                             sum((x["_b"][0] for x in both), [])), 3) if both else None,
                          agreement_supported=round(agreement(sum((x["_a"][1] for x in both), []),
                                                              sum((x["_b"][1] for x in both), [])), 3) if both else None)
        for x in rows:
            x.pop("_a", None), x.pop("_b", None)
    res["annotation"] = dict(summary=per, lectures=rows)
    # --- teacher ratings
    key_path = EVAL / "rating" / "rating_key.json"
    if key_path.exists():
        key = {str(it["item"]): it for it in json.loads(key_path.read_text(encoding="utf-8"))}
        scores: dict[str, dict[tuple[str, str], float]] = {}
        for f in sorted((EVAL / "rating" / "returned").glob("*.csv")):
            with f.open(encoding="utf-8-sig") as fh:
                for row in csv.DictReader(fh):
                    if row["criterion"] != "comment" and row["score"].strip():
                        scores.setdefault(row["rater"], {})[(row["item"], row["criterion"])] = float(row["score"])
        raters = sorted(scores)
        if raters:
            out = dict(raters=raters, items=len(key))
            for c, _ in CRITERIA:
                m = np.array([[scores[r].get((i, c), np.nan) for r in raters] for i in key])
                cell = dict(mean=round(float(np.nanmean(m)), 2), sd=round(float(np.nanstd(m, ddof=1)), 2),
                            icc=round(icc2(m), 3) if len(raters) > 1 else None)
                for g in sorted({group_of(it["group"]) for it in key.values()}):
                    gm = np.array([[scores[r].get((i, c), np.nan) for r in raters] for i in key
                                   if group_of(key[i]["group"]) == g])
                    cell[g] = round(float(np.nanmean(gm)), 2)
                out[c] = cell
            res["ratings"] = out
    EVAL.mkdir(exist_ok=True)
    (EVAL / "results.json").write_text(json.dumps(res, indent=1, ensure_ascii=False, default=float), encoding="utf-8")
    print(json.dumps({k: v for k, v in res.items() if k != "annotation"} | {"annotation": per}, indent=1, default=float))
    print(f"\n-> {EVAL / 'results.json'}")


if __name__ == "__main__":
    {"sheets": make_sheets, "score": score}[sys.argv[1] if len(sys.argv) > 1 else "score"]()
