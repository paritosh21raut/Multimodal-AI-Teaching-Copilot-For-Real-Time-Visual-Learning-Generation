from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation

PROJECT_ROOT = Path(__file__).resolve().parents[1]

if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.slides.slide_renderer import slide_renderer


def create_slide():
    presentation = Presentation()

    # Use the same layout used by PPTManager.
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[1]
    )

    return presentation, slide


def test_visual(
    visual_type,
    title,
    bullets,
    visual_spec,
):
    presentation, slide = create_slide()

    slide_renderer.render_content_slide(
        slide=slide,
        title=title,
        bullets=bullets,
        visual_type=visual_type,
        visual_spec=visual_spec,
    )

    output = (
        PROJECT_ROOT
        / "outputs"
        / f"test_{visual_type}.pptx"
    )

    presentation.save(output)

    print(
        f"[PASS] {visual_type:<20} -> {output}"
    )


def main():

    print("=" * 70)
    print("VISUAL RENDERER TEST")
    print("=" * 70)

    test_visual(
        "none",
        "Introduction to Microcontrollers",
        [
            "Small computer on a single chip.",
            "Contains processor, memory and I/O.",
            "Used in embedded systems.",
        ],
        {},
    )

    test_visual(
        "diagram",
        "Microcontroller Architecture",
        [
            "Processor controls the system.",
            "Memory stores programs and data.",
            "I/O connects external devices.",
        ],
        {
            "center": "Microcontroller",
            "components": [
                "CPU",
                "Memory",
                "I/O",
            ],
        },
    )

    test_visual(
        "flowchart",
        "How a Microcontroller Works",
        [
            "Receives input from sensors.",
            "Processes the input.",
            "Produces an output.",
        ],
        {
            "nodes": [
                "Input",
                "Process",
                "Decision",
                "Output",
            ],
        },
    )

    test_visual(
        "comparison_table",
        "8-bit vs 16-bit vs 32-bit",
        [
            "Microcontrollers differ by data width.",
            "Higher width can support larger data operations.",
        ],
        {
            "columns": [
                "Feature",
                "8-bit",
                "16-bit",
                "32-bit",
            ],
            "rows": [
                {
                    "label": "Data width",
                    "values": [
                        "8 bits",
                        "16 bits",
                        "32 bits",
                    ],
                },
                {
                    "label": "Typical use",
                    "values": [
                        "Simple",
                        "Medium",
                        "Advanced",
                    ],
                },
            ],
        },
    )

    test_visual(
        "chart",
        "Example Performance Data",
        [
            "Chart is used only for supplied numerical data.",
        ],
        {
            "chart_type": "bar",
            "x": [
                "8-bit",
                "16-bit",
                "32-bit",
            ],
            "y": [
                8,
                16,
                32,
            ],
        },
    )

    test_visual(
        "hierarchy",
        "Microcontroller Classification",
        [
            "Microcontrollers can be grouped by architecture.",
        ],
        {
            "root": "Microcontrollers",
            "levels": [
                {
                    "name": "Data Width",
                    "children": [
                        "8-bit",
                        "16-bit",
                        "32-bit",
                    ],
                }
            ],
        },
    )

    test_visual(
        "example_grid",
        "Applications of Microcontrollers",
        [
            "Microcontrollers are used in many embedded devices.",
        ],
        {
            "examples": [
                "Air Conditioner",
                "Robotics",
                "Washing Machine",
                "Automotive",
                "IoT",
                "Consumer Electronics",
            ]
        },
    )

    test_visual(
        "timeline",
        "Evolution of a System",
        [
            "A timeline is useful for ordered events.",
        ],
        {
            "events": [
                "Design",
                "Prototype",
                "Testing",
                "Deployment",
            ],
        },
    )

    test_visual(
        "formula",
        "Basic Formula",
        [
            "Formula is shown separately when appropriate.",
        ],
        {
            "formula": "V = I × R",
        },
    )

    test_visual(
        "concept_map",
        "Microcontroller Applications",
        [
            "Microcontrollers connect computing with applications.",
        ],
        {
            "center": "Microcontroller",
            "concepts": [
                "Robotics",
                "IoT",
                "Automotive",
                "Automation",
            ],
        },
    )

    print("=" * 70)
    print("VISUAL RENDERER TEST PASSED")
    print("=" * 70)


if __name__ == "__main__":
    main()