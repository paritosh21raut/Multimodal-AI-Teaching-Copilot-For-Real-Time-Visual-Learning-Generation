"""
Slide Renderer - Comprehensive Tests
"""

from __future__ import annotations

import sys
import os
import pytest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '../..')))

from pptx import Presentation
from pptx.util import Inches

from app.presentation.models.presentation_models import (
    SlidePlan,
    LayoutPlan,
    LayoutFamily,
    ContentBlock,
    ContentBlockType,
    RepresentationType,
    RepresentationDecision,
)
from app.presentation.renderer.slide_renderer import SlideRenderer
from app.presentation.layout.layout_engine import LayoutEngine
from app.presentation.design.design_system import design_system


def create_test_plan(layout_family=LayoutFamily.TEXT_LEFT_VISUAL_RIGHT):
    """Helper to create test plan"""
    return SlidePlan(
        slide_id="test",
        focal_message="TCP provides reliable delivery",
        content_blocks=[
            ContentBlock(block_type=ContentBlockType.KEY_CLAIM, text="TCP is connection-oriented"),
            ContentBlock(block_type=ContentBlockType.EXPLANATION, text="TCP PROVIDES reliability"),
        ],
        representation=RepresentationDecision(
            representation_type=RepresentationType.EXPLANATION,
            confidence=0.8,
            reason="Test",
        ),
        layout_family=layout_family,
        density=0.5,
    )


# EASY TESTS

def test_renderer_initializes():
    """Easy: Renderer initializes"""
    renderer = SlideRenderer()
    assert renderer is not None


def test_render_slide_creates_shapes():
    """Easy: Rendering creates shapes"""
    renderer = SlideRenderer()
    engine = LayoutEngine()
    
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    
    plan = create_test_plan()
    layout = engine.create_layout(plan)
    
    renderer.render_slide(slide, plan, layout)
    
    assert len(slide.shapes) >= 2  # Title + content


def test_render_hero_definition():
    """Easy: Hero definition renders"""
    renderer = SlideRenderer()
    engine = LayoutEngine()
    
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    
    plan = create_test_plan(LayoutFamily.HERO_DEFINITION)
    layout = engine.create_layout(plan)
    
    renderer.render_slide(slide, plan, layout)
    
    assert len(slide.shapes) >= 3  # Title + definition + accent line


# MEDIUM TESTS

def test_all_layout_families_render():
    """Medium: All layout families render without crash"""
    renderer = SlideRenderer()
    engine = LayoutEngine()
    
    for family in LayoutFamily:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        
        plan = create_test_plan(family)
        layout = engine.create_layout(plan)
        
        # Should not crash
        renderer.render_slide(slide, plan, layout)
        
        assert len(slide.shapes) >= 1


def test_render_comparison_with_table():
    """Medium: Comparison with table spec"""
    renderer = SlideRenderer()
    engine = LayoutEngine()
    
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    
    plan = create_test_plan(LayoutFamily.FULL_WIDTH_COMPARISON)
    layout = engine.create_layout(plan)
    
    visual_spec = {
        "type": "comparison_table",
        "columns": ["TCP", "UDP"],
        "rows": [
            {"label": "Connection", "values": ["Oriented", "Connectionless"]},
            {"label": "Reliability", "values": ["Reliable", "Unreliable"]},
        ],
    }
    
    renderer.render_slide(slide, plan, layout, visual_spec)
    
    assert len(slide.shapes) >= 2


def test_render_process_with_flowchart():
    """Medium: Process with flowchart spec"""
    renderer = SlideRenderer()
    engine = LayoutEngine()
    
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    
    plan = create_test_plan(LayoutFamily.FULL_WIDTH_PROCESS)
    layout = engine.create_layout(plan)
    
    visual_spec = {
        "type": "flowchart",
        "nodes": ["Step 1", "Step 2", "Step 3"],
        "edges": [["Step 1", "Step 2"], ["Step 2", "Step 3"]],
    }
    
    renderer.render_slide(slide, plan, layout, visual_spec)
    
    assert len(slide.shapes) >= 4  # Title + 3 nodes


# HARD TESTS

def test_render_with_dense_content():
    """Hard: Dense content renders without overflow"""
    renderer = SlideRenderer()
    engine = LayoutEngine()
    
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    
    plan = SlidePlan(
        focal_message="Test",
        content_blocks=[
            ContentBlock(block_type=ContentBlockType.EXPLANATION, text=f"Block {i}")
            for i in range(10)
        ],
        layout_family=LayoutFamily.TEXT_LEFT_VISUAL_RIGHT,
        density=0.9,
    )
    layout = engine.create_layout(plan)
    
    renderer.render_slide(slide, plan, layout)
    
    assert len(slide.shapes) >= 2


def test_render_empty_plan():
    """Hard: Empty plan doesn't crash"""
    renderer = SlideRenderer()
    engine = LayoutEngine()
    
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    
    plan = SlidePlan(
        focal_message="",
        content_blocks=[],
        layout_family=LayoutFamily.TEXT_LEFT_VISUAL_RIGHT,
        density=0.5,
    )
    layout = engine.create_layout(plan)
    
    renderer.render_slide(slide, plan, layout)
    
    assert len(slide.shapes) >= 1


def test_render_all_representations():
    """Hard: All representation types render"""
    renderer = SlideRenderer()
    engine = LayoutEngine()
    
    for rep_type in [RepresentationType.DEFINITION, RepresentationType.COMPARISON, 
                     RepresentationType.FLOWCHART, RepresentationType.NUMBER_STATISTIC]:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        
        plan = SlidePlan(
            focal_message="Test",
            content_blocks=[ContentBlock(block_type=ContentBlockType.KEY_CLAIM, text="Test")],
            representation=RepresentationDecision(
                representation_type=rep_type,
                confidence=0.8,
                reason="Test",
            ),
            layout_family=LayoutFamily.TEXT_LEFT_VISUAL_RIGHT,
            density=0.5,
        )
        layout = engine.create_layout(plan)
        
        renderer.render_slide(slide, plan, layout)
        
        assert len(slide.shapes) >= 1


# GENERIC TESTS

def test_renderer_domain_agnostic():
    """Generic: Renderer works regardless of domain"""
    renderer = SlideRenderer()
    engine = LayoutEngine()
    
    domains = {
        "networking": "TCP provides reliability",
        "biology": "Photosynthesis produces glucose",
        "physics": "Force causes acceleration",
    }
    
    for domain, claim in domains.items():
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        
        plan = SlidePlan(
            focal_message=claim,
            content_blocks=[ContentBlock(block_type=ContentBlockType.KEY_CLAIM, text=claim)],
            layout_family=LayoutFamily.HERO_DEFINITION,
            density=0.5,
        )
        layout = engine.create_layout(plan)
        
        renderer.render_slide(slide, plan, layout)
        
        assert len(slide.shapes) >= 1