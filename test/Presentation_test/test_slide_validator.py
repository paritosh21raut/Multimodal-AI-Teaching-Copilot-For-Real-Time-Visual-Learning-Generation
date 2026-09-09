"""
Slide Validator - Comprehensive Tests
"""

from __future__ import annotations

import sys
import os
import pytest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '../..')))

from pptx import Presentation
from pptx.util import Inches, Pt

from app.presentation.models.presentation_models import (
    SlideValidationResult,
    SlidePlan,
    ContentBlock,
    ContentBlockType,
    RepresentationType,
    RepresentationDecision,
    LayoutFamily,
)
from app.presentation.qa.slide_validator import SlideValidator
from app.presentation.renderer.slide_renderer import SlideRenderer
from app.presentation.layout.layout_engine import LayoutEngine


def create_test_plan():
    return SlidePlan(
        focal_message="TCP provides reliable delivery",
        content_blocks=[
            ContentBlock(block_type=ContentBlockType.KEY_CLAIM, text="TCP is connection-oriented"),
        ],
        layout_family=LayoutFamily.HERO_DEFINITION,
        density=0.5,
    )


# EASY TESTS

def test_validator_initializes():
    """Easy: Validator initializes"""
    validator = SlideValidator()
    assert validator is not None


def test_valid_rendered_slide():
    """Easy: Properly rendered slide passes validation"""
    validator = SlideValidator()
    renderer = SlideRenderer()
    engine = LayoutEngine()
    
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    
    plan = create_test_plan()
    layout = engine.create_layout(plan)
    renderer.render_slide(slide, plan, layout)
    
    result = validator.validate_slide(slide, plan)
    
    assert result is not None
    assert result.severity in ["none", "warning"]


def test_empty_slide_detected():
    """Easy: Empty slide detected"""
    validator = SlideValidator()
    
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    
    # Remove any default shapes
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    
    result = validator.validate_slide(slide)
    
    assert not result.is_valid
    assert result.severity == "severe"


# MEDIUM TESTS

def test_small_font_detected():
    """Medium: Small font detected"""
    validator = SlideValidator(min_font_size=12)
    
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Add shape with tiny font
    shape = slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(2), Inches(0.5)
    )
    shape.text_frame.text = "Tiny text"
    shape.text_frame.paragraphs[0].font.size = Pt(6)  # Below minimum
    
    result = validator.validate_slide(slide)
    
    assert result.severity in ["warning", "severe"]
    assert len(result.issues) >= 1


def test_out_of_bounds_detected():
    """Medium: Out of bounds shape detected"""
    validator = SlideValidator()
    
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Add shape outside slide bounds
    shape = slide.shapes.add_textbox(
        Inches(11), Inches(1), Inches(2), Inches(0.5)  # Left > 10" slide width
    )
    shape.text_frame.text = "Out of bounds"
    
    result = validator.validate_slide(slide)
    
    assert result.severity in ["warning", "severe"]


def test_too_many_shapes_detected():
    """Medium: Too many shapes flagged"""
    validator = SlideValidator(max_shapes=5)
    
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    for i in range(10):
        shape = slide.shapes.add_textbox(
            Inches(1), Inches(i * 0.5), Inches(2), Inches(0.3)
        )
        shape.text_frame.text = f"Shape {i}"
    
    result = validator.validate_slide(slide)
    
    assert len(result.issues) >= 1


# HARD TESTS

def test_shape_overlap_detected():
    """Hard: Overlapping shapes detected"""
    validator = SlideValidator()
    
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Two overlapping shapes
    shape1 = slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(2), Inches(1)
    )
    shape1.text_frame.text = "Shape 1"
    
    shape2 = slide.shapes.add_textbox(
        Inches(1.5), Inches(1.5), Inches(2), Inches(1)
    )
    shape2.text_frame.text = "Shape 2"
    
    result = validator.validate_slide(slide)
    
    assert len(result.issues) >= 1


def test_all_rendered_layouts_pass_validation():
    """Hard: All rendered layouts pass validation"""
    validator = SlideValidator()
    renderer = SlideRenderer()
    engine = LayoutEngine()
    
    for family in LayoutFamily:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        
        plan = create_test_plan()
        plan.layout_family = family
        
        layout = engine.create_layout(plan)
        renderer.render_slide(slide, plan, layout)
        
        result = validator.validate_slide(slide, plan)
        
        # Should not be severe
        assert result.severity != "severe", f"Severe issue for {family}: {result.issues}"


# GENERIC TESTS

def test_validator_generic():
    """Generic: Validator works regardless of content"""
    validator = SlideValidator()
    renderer = SlideRenderer()
    engine = LayoutEngine()
    
    domains = {
        "networking": "TCP provides reliability",
        "biology": "Photosynthesis produces glucose",
        "physics": "Force causes acceleration",
    }
    
    for domain, claim in domains.items():
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        
        plan = SlidePlan(
            focal_message=claim,
            content_blocks=[ContentBlock(block_type=ContentBlockType.KEY_CLAIM, text=claim)],
            layout_family=LayoutFamily.HERO_DEFINITION,
            density=0.5,
        )
        layout = engine.create_layout(plan)
        renderer.render_slide(slide, plan, layout)
        
        result = validator.validate_slide(slide, plan)
        
        assert result.severity != "severe", f"Failed for {domain}"