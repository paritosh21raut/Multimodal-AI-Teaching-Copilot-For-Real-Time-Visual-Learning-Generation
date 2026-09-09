"""
Phase 11: Renderer - Production Readiness Tests
"""

from __future__ import annotations

import sys
import os
import pytest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '../..')))

from pptx import Presentation

from app.presentation.renderer.slide_renderer import SlideRenderer
from app.presentation.layout.layout_engine import LayoutEngine
from app.presentation.models.presentation_models import (
    SlidePlan, LayoutFamily, ContentBlock, ContentBlockType,
    RepresentationType, RepresentationDecision,
)
from app.presentation.qa.slide_validator import SlideValidator


def render_and_validate(family):
    renderer = SlideRenderer()
    engine = LayoutEngine()
    validator = SlideValidator()
    
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    
    plan = SlidePlan(
        focal_message="TCP provides reliable delivery",
        content_blocks=[
            ContentBlock(block_type=ContentBlockType.KEY_CLAIM, text="TCP is connection-oriented"),
        ],
        representation=RepresentationDecision(
            representation_type=RepresentationType.EXPLANATION,
            confidence=0.8,
            reason="Test",
        ),
        layout_family=family,
        density=0.5,
    )
    
    layout = engine.create_layout(plan)
    renderer.render_slide(slide, plan, layout)
    return validator.validate_slide(slide, plan)


def test_hero_definition_renders():
    result = render_and_validate(LayoutFamily.HERO_DEFINITION)
    assert result.severity != "severe"


def test_all_layouts_render_without_severe():
    for family in LayoutFamily:
        result = render_and_validate(family)
        assert result.severity != "severe", f"Severe for {family}: {result.issues}"


def test_shapes_created():
    renderer = SlideRenderer()
    engine = LayoutEngine()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    plan = SlidePlan(
        focal_message="Test", content_blocks=[
            ContentBlock(block_type=ContentBlockType.KEY_CLAIM, text="Test claim"),
        ],
        layout_family=LayoutFamily.TEXT_LEFT_VISUAL_RIGHT, density=0.5,
    )
    layout = engine.create_layout(plan)
    renderer.render_slide(slide, plan, layout)
    assert len(slide.shapes) >= 2


def test_empty_plan_no_crash():
    renderer = SlideRenderer()
    engine = LayoutEngine()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    plan = SlidePlan(focal_message="", content_blocks=[], layout_family=LayoutFamily.TEXT_LEFT_VISUAL_RIGHT, density=0.5)
    layout = engine.create_layout(plan)
    renderer.render_slide(slide, plan, layout)
    assert len(slide.shapes) >= 1