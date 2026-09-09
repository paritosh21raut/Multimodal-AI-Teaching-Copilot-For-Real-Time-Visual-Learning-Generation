"""
Phase 12-20: QA, Presentation State, LLM Provider, Dashboard, Integration
Production Readiness Tests
"""

from __future__ import annotations

import sys
import os
import pytest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '../..')))

from pptx import Presentation
from pptx.util import Inches, Pt

from app.presentation.qa.slide_validator import SlideValidator
from app.presentation.llm.llm_router import LLMRouter
from app.presentation.llm.deterministic_provider import DeterministicProvider
from app.presentation.dashboard.dashboard_builder import DashboardBuilder
from app.presentation.bridge import PresentationBridge
from app.presentation.integration import PresentationIntelligence


# ============================================================
# PHASE 12: QA
# ============================================================

def test_qa_empty_slide_detected():
    validator = SlideValidator()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    result = validator.validate_slide(slide)
    assert not result.is_valid
    assert result.severity == "severe"


def test_qa_out_of_bounds_detected():
    validator = SlideValidator()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(11), Inches(1), Inches(2), Inches(0.5))
    shape.text_frame.text = "Out of bounds"
    result = validator.validate_slide(slide)
    assert len(result.issues) >= 1


# ============================================================
# PHASE 13: LLM Provider
# ============================================================

def test_deterministic_provider_always_works():
    provider = DeterministicProvider()
    result = provider.generate_structured(prompt="TOPIC: TCP\n- TCP PROVIDES reliability")
    assert result is not None
    assert "title" in result


def test_router_never_returns_none():
    router = LLMRouter()
    result = router.generate_structured(prompt="")
    assert result is not None


def test_router_tracks_providers():
    router = LLMRouter()
    status = router.get_provider_status()
    assert "groq" in status
    assert "deterministic" in status


# ============================================================
# PHASE 14: Dashboard
# ============================================================

def test_dashboard_teacher_student_separate():
    builder = DashboardBuilder()
    builder.update_teacher_view(topic="TCP")
    builder.update_student_view(slide_title="TCP Slide")
    
    teacher = builder.get_teacher_view()
    student = builder.get_student_view()
    
    assert teacher is not None
    assert student is not None
    assert teacher != student


def test_dashboard_mode_switching():
    builder = DashboardBuilder()
    builder.set_mode("teacher")
    assert builder.get_state().mode == "teacher"
    builder.set_mode("student")
    assert builder.get_state().mode == "student"


# ============================================================
# PHASE 15-20: Integration
# ============================================================

def test_bridge_initializes():
    bridge = PresentationBridge()
    assert bridge is not None


def test_presentation_intelligence_initializes():
    pi = PresentationIntelligence()
    assert pi is not None
    state = pi.get_state()
    assert "slide_decision" in state
    assert "llm_providers" in state


def test_full_pipeline_initialization():
    """Every subsystem initializes without crash"""
    from app.presentation.slide_decision.slide_decision_engine import SlideDecisionEngine
    from app.presentation.information_selection.information_selector import InformationSelector
    from app.presentation.representation.representation_engine import RepresentationEngine
    from app.presentation.planner.presentation_planner import PresentationPlanner
    from app.presentation.layout.layout_engine import LayoutEngine
    from app.presentation.visuals.structured_visual_engine import StructuredVisualEngine
    from app.presentation.visual_assets.visual_asset_engine import VisualAssetEngine
    from app.presentation.renderer.slide_renderer import SlideRenderer
    from app.presentation.qa.slide_validator import SlideValidator
    from app.presentation.llm.llm_router import LLMRouter
    from app.presentation.dashboard.dashboard_builder import DashboardBuilder
    
    assert SlideDecisionEngine() is not None
    assert InformationSelector() is not None
    assert RepresentationEngine() is not None
    assert PresentationPlanner() is not None
    assert LayoutEngine() is not None
    assert StructuredVisualEngine() is not None
    assert VisualAssetEngine() is not None
    assert SlideRenderer() is not None
    assert SlideValidator() is not None
    assert LLMRouter() is not None
    assert DashboardBuilder() is not None
    