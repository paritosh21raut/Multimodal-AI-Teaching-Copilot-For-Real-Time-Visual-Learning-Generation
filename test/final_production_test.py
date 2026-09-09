"""
FINAL PRODUCTION TEST

Complete system with all fixes:
- Semantic relation correctness
- Development coverage
- Importance factors
- Slide decision with dedup
- Information selection
- Representation
- Planner
- Layout
- Renderer
- QA

Speak for 60 seconds. System should generate a PPT.
"""

from __future__ import annotations

import sys
import os
import time
import threading

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from app.audio.audio_stream import AudioStream
from app.audio.audio_worker import AudioWorker
from app.speech.transcript_intelligence import transcript_intelligence
from app.topics.topic_intelligence import TopicIntelligence
from app.presentation.bridge import PresentationBridge
from app.semantic.semantic_pipeline import SemanticPipeline
from app.development.development_tracker import DevelopmentTracker
from app.importance.importance_scorer import ImportanceScorer

from pptx import Presentation

from app.presentation.renderer.slide_renderer import SlideRenderer
from app.presentation.layout.layout_engine import LayoutEngine
from app.presentation.qa.slide_validator import SlideValidator
from app.presentation.design.design_system import design_system


class ProductionTester:
    """Complete production test"""
    
    def __init__(self, duration=60):
        self.duration = duration
        
        self.audio_stream = AudioStream()
        self.audio_worker = AudioWorker(
            audio_queue=self.audio_stream.get_queue(),
            on_preview_transcript=self._on_preview,
            on_final_transcript=self._on_final,
        )
        
        self.transcript_intelligence = transcript_intelligence
        self.topic_intelligence = TopicIntelligence()
        self.semantic_pipeline = SemanticPipeline()
        self.development_tracker = DevelopmentTracker()
        self.importance_scorer = ImportanceScorer(self.development_tracker)
        self.presentation_bridge = PresentationBridge()
        
        self.slide_renderer = SlideRenderer()
        self.layout_engine = LayoutEngine()
        self.slide_validator = SlideValidator()
        
        self.prs = Presentation()
        self.slide_count = 0
        self.ppt_path = f"outputs/presentations/Production_Test_{int(time.time())}.pptx"
        
        self.current_topic = None
        self.current_embedding = None
        self.context = ""
        
        self.final_count = 0
        self.preview_count = 0
        self.full_transcript = ""
        self.topics = []
        
        os.makedirs("outputs/presentations", exist_ok=True)
    
    def _on_preview(self, text, authoritative):
        self.preview_count += 1
    
    def _on_final(self, segment_text, full_transcript, segment_id):
        self.final_count += 1
        self.full_transcript = full_transcript
        
        refined = self.transcript_intelligence.refine(segment_text, context=self.context)
        refined_text = refined.refined_text.strip()
        
        if not refined_text:
            return
        
        topic_decision = self.topic_intelligence.process(
            latest_text=refined_text,
            rolling_context=self.context,
            current_topic=self.current_topic,
            current_embedding=self.current_embedding,
        )
        
        if not topic_decision.is_relevant:
            return
        
        self.current_topic = topic_decision.topic
        self.current_embedding = topic_decision.embedding
        self.context = f"{self.context} {refined_text}".strip()
        
        if topic_decision.topic not in self.topics:
            self.topics.append(topic_decision.topic)
        
        semantic_result = self.semantic_pipeline.process_transcript(
            transcript_text=refined_text,
            chunk_id=f"chunk_{self.final_count}",
            lecture_id="production_test",
            topic_path=str(topic_decision.topic),
            generate_embeddings=False,
        )
        
        if not semantic_result or not semantic_result.get("frame"):
            return
        
        frame = semantic_result["frame"]
        
        self.development_tracker.process_frame(frame, chunk_id=f"chunk_{self.final_count}")
        
        important = self.importance_scorer.get_top_important(limit=5)
        important_names = [c.canonical_name for c in important]
        
        result = self.presentation_bridge.process_transcript(
            frame=frame,
            topic_changed=topic_decision.is_new_topic,
            current_topic=str(topic_decision.topic),
            important_concepts=important_names,
            authoritative_transcript=self.full_transcript,
            chunk_id=f"chunk_{self.final_count}",
        )
        
        if result and result.get("plan") and result.get("slide_created"):
            plan = result["plan"]
            layout = result.get("layout")
            visual_spec = result.get("visual_spec")
            
            # Render slide
            slide = self.prs.slides.add_slide(self.prs.slide_layouts[1])
            self.slide_renderer.render_slide(slide, plan, layout, visual_spec)
            
            # QA
            validation = self.slide_validator.validate_slide(slide, plan)
            
            if validation.severity == "severe":
                print(f"  ❌ Slide rejected by QA: {validation.issues}")
            else:
                self.slide_count += 1
                self.prs.save(self.ppt_path)
                print(f"  ✅ Slide {self.slide_count}: {plan.focal_message[:50]}")
        
        print(f"[Topic] {topic_decision.topic[:60]}")
        print(f"[Concepts] {len(frame.concepts)}")
        print(f"[Important] {important_names[:3]}")
        print(f"[Action] {result.get('action', 'no_action') if result else 'none'}")
    
    def run(self):
        print("="*70)
        print("FINAL PRODUCTION TEST")
        print("="*70)
        print("Speak about TCP/UDP networking concepts")
        print("="*70)
        
        worker_thread = threading.Thread(target=self.audio_worker.run, daemon=True)
        worker_thread.start()
        time.sleep(2)
        
        stream_thread = threading.Thread(target=self.audio_stream.start, daemon=True)
        stream_thread.start()
        
        try:
            start = time.time()
            while time.time() - start < self.duration:
                if not worker_thread.is_alive():
                    break
                time.sleep(1)
        except KeyboardInterrupt:
            pass
        finally:
            self.audio_stream.stop()
            self.audio_worker.stop()
            worker_thread.join(timeout=15)
            time.sleep(3)
        
        if self.slide_count > 0:
            self.prs.save(self.ppt_path)
        
        print("\n" + "="*70)
        print("RESULTS")
        print("="*70)
        print(f"Final transcripts: {self.final_count}")
        print(f"Previews: {self.preview_count}")
        print(f"Slides: {self.slide_count}")
        print(f"PPT: {self.ppt_path}")
        print(f"Topics: {self.topics}")
        
        if os.path.exists(self.ppt_path):
            print(f"✅ PPT saved: {os.path.abspath(self.ppt_path)}")


if __name__ == "__main__":
    tester = ProductionTester(duration=60)
    tester.run()