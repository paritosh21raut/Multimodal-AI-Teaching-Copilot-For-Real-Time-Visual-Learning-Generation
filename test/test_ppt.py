from pathlib import Path
from pptx import Presentation

# Create output directory
output_dir = Path("outputs/presentations")

if not output_dir.exists():
    output_dir.mkdir(parents=True)

# Create presentation
prs = Presentation()

# Add one slide
layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(layout)

slide.shapes.title.text = "Lecture Test"

slide.placeholders[1].text = "PowerPoint generation is working successfully."

# Save
output_file = output_dir / "Lecture_Test.pptx"
prs.save(output_file)

print(f"Presentation created successfully:\n{output_file.resolve()}")