"""My notes beyond PDFs (F-010b §5): images, text, Word / PowerPoint (real Office: slow), web links, and a notes page
on the projector — followed by page turns, never moved by the lecture, private pages stay private."""
import io
import shutil
from pathlib import Path

import httpx
import pytest

from copilot.core.events import NotesProjected
from copilot.display.hub import DisplayHub
from copilot.display.server import DisplayServer
from copilot.notes.convert import BadNotes, clean_url, image_pdf, kind_of, office_pdf, text_page, to_pdf
from copilot.notes.library import NotesLibrary

from tests.integration.test_display_hub import free_port
from tests.integration.test_share_access import TUNNEL
from tests.unit.test_notes import PDF, _live, _service, cmd


def png(w=300, h=200) -> bytes:
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGBA", (w, h), (20, 120, 110, 255)).save(out, format="PNG")
    return out.getvalue()


def test_kinds_by_extension():
    assert [kind_of(n) for n in ("a.PDF", "b.docx", "c.ppt", "d.md", "e.jpeg", "f.exe", "noext")] == \
        ["pdf", "word", "slides", "text", "image", None, None]


async def test_an_image_becomes_a_one_page_pdf_without_text(tmp_path):
    pdf, kind = await to_pdf(png(), "Leaf diagram.png")
    assert kind == "image" and pdf.startswith(b"%PDF")
    doc = NotesLibrary(tmp_path).add(pdf, "Leaf diagram.png", kind)
    assert doc.pages == 1 and not doc.has_text and doc.kind == "image"
    with pytest.raises(BadNotes):
        image_pdf(b"not an image")
    with pytest.raises(BadNotes):
        await to_pdf(b"MZ", "virus.exe")


def test_a_markdown_page_keeps_headings_and_bullets():
    page = text_page("Plan", "# Week 1\n- light <b>\n\nplain line", markdown=True)
    assert "<h2>Week 1</h2>" in page and "• light &lt;b&gt;" in page and "<p>plain line</p>" in page


def test_only_web_links_are_taken():
    assert clean_url(" https://en.wikipedia.org/wiki/Leaf ") == "https://en.wikipedia.org/wiki/Leaf"
    for bad in ("file:///C:/x", "javascript:alert(1)", "wikipedia", ""):
        with pytest.raises(BadNotes):
            clean_url(bad)


async def test_a_page_on_the_projector_follows_page_turns_but_not_the_lecture(tmp_path):
    bus, svc, doc, states = await _service(tmp_path)
    shown: list[NotesProjected] = []

    async def grab(e):
        shown.append(e)
    bus.subscribe("t_shown", grab, [NotesProjected])
    await cmd(bus, "notes_open", id=doc.id)
    await cmd(bus, "notes_project", on=True)
    assert shown[-1].on and shown[-1].doc_id == doc.id and shown[-1].page == 1 and states[-1].projecting
    await cmd(bus, "notes_page", page=4)
    assert shown[-1].page == 4
    await _live(bus, "s2", "Respiration", "Respiration releases energy from glucose, respiration")
    assert states[-1].page == 4 and shown[-1].page == 4       # the class is reading: the lecture does not turn it
    await cmd(bus, "notes_project", on=False)
    assert not shown[-1].on and shown[-1].doc_id == ""
    await cmd(bus, "notes_project", on=True)
    await cmd(bus, "notes_remove", id=doc.id)                 # removed while shown: off the projector
    assert not shown[-1].on and not states[-1].projecting
    await bus.close()


async def test_only_the_projected_page_reaches_students(tmp_path):
    from copilot.core.bus import EventBus

    library = NotesLibrary(tmp_path / "notes")
    doc = library.add(PDF.read_bytes(), "photosynthesis_notes.pdf")
    bus = EventBus()
    hub = DisplayHub(bus)
    hub.attach()
    viewer = hub.connect("viewer")
    await viewer.next_batch()
    server = DisplayServer(hub, port=free_port(), control_key="k" * 32, notes=library)
    await server.start()
    try:
        async with httpx.AsyncClient(base_url=f"http://127.0.0.1:{server.port}", headers=TUNNEL) as student:
            assert (await student.get(f"/api/notes/shown/{doc.id}/2")).status_code == 403   # nothing shown
            await bus.publish(NotesProjected(on=True, doc_id=doc.id, page=2, name=doc.name))
            await bus.drain()
            assert any(m["type"] == "notes_shown" and m["page"] == 2 for m in await viewer.next_batch())
            r = await student.get(f"/api/notes/shown/{doc.id}/2")
            assert r.status_code == 200 and r.headers["content-type"] == "image/png"
            assert (await student.get(f"/api/notes/shown/{doc.id}/3")).status_code == 403  # another page: private
            assert (await student.get(f"/api/notes/{doc.id}/page/2")).status_code == 403
            assert (await student.post("/api/notes/link", json={"url": "https://example.org"})).status_code == 403
        async with httpx.AsyncClient(base_url=f"http://127.0.0.1:{server.port}") as teacher:
            r = await teacher.post("/api/notes", content=png(), headers={"X-File-Name": "leaf.png"})
            assert r.status_code == 200 and r.json()["pages"] == 1
            r = await teacher.post("/api/notes", content=b"MZ", headers={"X-File-Name": "tool.exe"})
            assert r.status_code == 415
            r = await teacher.post("/api/notes/link", json={"url": "notalink"})
            assert r.status_code == 422
    finally:
        await server.stop()
        await bus.close()


OFFICE = shutil.which("powershell") and Path(r"C:\Program Files\Microsoft Office\root\Office16\WINWORD.EXE").exists()


@pytest.mark.slow
@pytest.mark.skipif(not OFFICE, reason="needs Microsoft Word + PowerPoint")
def test_word_and_powerpoint_files_become_pdfs_through_office(tmp_path):
    from pptx import Presentation

    (tmp_path / "n.rtf").write_text(r"{\rtf1\ansi{\b Photosynthesis}\par Green plants make glucose from carbon dioxide "
                                    r"and water using sunlight.\par}", encoding="ascii")
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Respiration"
    s.placeholders[1].text = "Glucose is broken down to release energy."
    p.save(tmp_path / "n.pptx")
    lib = NotesLibrary(tmp_path / "notes")
    for name, kind, words in (("n.rtf", "word", "glucose from carbon dioxide"), ("n.pptx", "slides", "release energy")):
        pdf = office_pdf((tmp_path / name).read_bytes(), name, kind)
        doc = lib.add(pdf, name, kind)
        assert doc.pages >= 1 and words in " ".join(lib.page_texts(doc.id)).lower()
