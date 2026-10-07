"""Lecture materials + past lectures in real Edge (F-010): the teacher ticks what to make in /control, Create makes
real files (Edge prints the PDFs, draws the PPTX), the summary slide reaches the projector, past lectures open in a
viewer, and students download the shared PDFs from their page. The LLM is a fake (0 tokens)."""
import asyncio
import io
import sys
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "tools"))
from display_harness import browser_page, display_harness  # noqa: E402

from copilot.core.events import TranscriptFinal, TranscriptSegment  # noqa: E402
from tests.unit.test_materials import PHOTO, SAID, FakeRouter, write_lecture  # noqa: E402

pytestmark = pytest.mark.browser
OUT = ROOT / "artifacts" / "materials"
NOTES_PDF = ROOT / "tests" / "fixtures" / "notes" / "photosynthesis_notes.pdf"


async def lecture_now(h):
    for s in PHOTO:
        await h.deck.add(s.model_copy())
    for lines in SAID.values():
        for line in lines:
            await h.bus.publish(TranscriptFinal(segment=TranscriptSegment(text=line, start=1, end=2)))
    await h.settle()


async def test_the_teacher_makes_materials_and_opens_a_past_lecture(tmp_path):
    OUT.mkdir(parents=True, exist_ok=True)
    await write_lecture(tmp_path / "sessions", "past", slides=PHOTO[1:4])
    async with display_harness(materials_dir=tmp_path, sessions_dir=tmp_path / "sessions", session_id="now",
                               materials_router=FakeRouter(), store=True) as h:
        await lecture_now(h)
        await h.bus.publish(__import__("copilot.core.events", fromlist=["CommandReceived"]).CommandReceived(
            command=__import__("copilot.core.events", fromlist=["Command"]).Command(kind="goto", args={"slide_id": "s2"})))
        await h.settle()
        async with browser_page(f"{h.url}/control", 1600, 1000) as control:
            await control.click(".tabs button[aria-label=Materials]")
            await control.wait_for_selector(".materials .kind")
            await control.screenshot(path=str(OUT / "materials_empty.png"))
            card = lambda label: f".kind:has(.kind-text b:text-is('{label}'))"  # noqa: E731  (the label, exactly)
            for kind in ("Summary slide", "Notes", "Assignment", "Slides"):
                await control.click(f"{card(kind)} .kind-head")
            placeholder = await control.get_attribute(f"{card('Notes')} input.name", "placeholder")
            assert placeholder.startswith("Photosynthesis - Notes - ") and placeholder.endswith(".pdf")
            await control.fill(f"{card('Notes')} input.name", "Week 1 notes")
            await control.fill(f"{card('Assignment')} .stepper input", "6")
            await control.press(f"{card('Assignment')} .stepper input", "Tab")
            await control.click(f"{card('Slides')} .segmented button:has-text('Dark')")
            await control.screenshot(path=str(OUT / "materials_ticked.png"))
            await control.click(".make-foot button.primary")
            await control.wait_for_function("document.querySelectorAll('.made.ready').length === 4", timeout=90000)
            await control.screenshot(path=str(OUT / "materials_made.png"))
            assert not control.errors, control.errors

            by = {m.kind: m for m in h.materials.store.items()}
            notes = h.materials.store.path(by["notes"])
            assert by["notes"].name == "Week 1 notes.pdf" and notes.read_bytes()[:5] == b"%PDF-"
            assert h.materials.store.path(by["assignment"]).read_bytes()[:5] == b"%PDF-" and by["assignment"].count == 6
            from pptx import Presentation

            prs = Presentation(str(h.materials.store.path(by["pptx"])))
            texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
            assert len(prs.slides) == 5 and "How it happens" in texts          # editable text, as on the slides
            assert h.deck.live.origin == "materials" and h.deck.live.title == "Summary"  # on the projector

            # a past lecture: the list, the viewer, its slides
            await control.click(".tabs button[aria-label=Lectures]")
            assert await control.get_attribute(".tabs button.on", "aria-label") == "Lectures"
            # the list now holds the running lecture too ("now", F-010b): open the past one
            await control.click(".lecture-row:not(.now) .lecture-open:has-text('Photosynthesis')")
            await control.wait_for_selector(".viewer .stage .slide")
            title_px = await control.eval_on_selector(".viewer .stage .slide-title", "e => parseFloat(getComputedStyle(e).fontSize)")
            assert title_px >= 50                                   # the slide as on the projector, not restyled
            await control.keyboard.press("ArrowRight")
            assert (await control.inner_text(".viewer-nav .count")).strip() == "2 / 3"
            await asyncio.sleep(0.5)
            await control.screenshot(path=str(OUT / "lecture_viewer.png"))
            await control.click(".viewer-head button.primary")   # Make materials from it
            # F-010b: From = Pick lectures, with that lecture picked
            await control.wait_for_selector(".materials .segmented button.on:has-text('Pick lectures')")
            await control.wait_for_selector(".materials .pick-wrap .menu-button:has-text('Photosynthesis')")
            assert await control.query_selector(".viewer") is None
            assert not control.errors, control.errors

        # the PDFs as a person sees them (looked at, not only generated)
        import pypdfium2

        for kind in ("notes", "assignment"):
            pdf = pypdfium2.PdfDocument(str(h.materials.store.path(by[kind])))
            for i in range(min(2, len(pdf))):
                pdf[i].render(scale=1.4).to_pil().save(OUT / f"{kind}_p{i + 1}.png")
            pdf.close()
        assert zipfile.is_zipfile(io.BytesIO(h.materials.store.path(by["pptx"]).read_bytes()))


async def test_students_download_the_shared_pdfs(tmp_path):
    async with display_harness(materials_dir=tmp_path, sessions_dir=tmp_path / "sessions", session_id="now",
                               materials_router=FakeRouter()) as h:
        await lecture_now(h)
        async with browser_page(f"{h.url}/view", 1280, 720) as view:
            assert await view.query_selector(".shared-files") is None   # nothing shared yet: no button
            from copilot.core.events import Command, CommandReceived

            await h.bus.publish(CommandReceived(command=Command(kind="materials_create", args={
                "items": [{"kind": "notes", "name": "Notes for class"}], "current": True})))
            await h.settle()
            assert await h.materials.wait_idle(60)
            await view.click(".shared-files > button", timeout=10000)
            link = await view.wait_for_selector(".shared-list a:has-text('Notes for class.pdf')")
            r = await view.request.get(h.url + await link.get_attribute("href"))
            assert r.status == 200 and (await r.body())[:5] == b"%PDF-"
            await view.screenshot(path=str(OUT / "student_materials.png"))
            assert not view.errors, view.errors


async def test_notes_menu_and_the_solid_moon(tmp_path):
    async with display_harness(notes_dir=tmp_path / "notes", store=True) as h:
        async with browser_page(f"{h.url}/control", 1600, 1000) as control:
            fill = await control.get_attribute(".dock button.theme svg.moon path", "fill")
            assert fill == "currentColor"                              # a solid crescent, not an outline
            await control.click(".tabs button[aria-label='My notes']")
            await control.wait_for_selector(".notes-empty button.primary")
            assert await control.query_selector(".notes-empty p") is None    # minimal: no paragraph
            await control.screenshot(path=str(OUT / "notes_empty_minimal.png"))
            async with control.expect_file_chooser() as chooser:
                await control.click(".notes-empty button.primary")
            await (await chooser.value).set_files(str(NOTES_PDF))
            await control.wait_for_selector(".doc-menu .menu-button")
            assert await control.query_selector("select") is None          # no native select
            await control.click(".doc-menu .menu-button")
            await control.wait_for_selector(".menu .menu-row.add:has-text('Add a file')")
            await control.wait_for_function("() => { const i = document.querySelector('.notes-page img');"
                                            " return !!i && i.complete && i.naturalWidth > 0; }")
            await asyncio.sleep(0.4)  # the menu's entrance
            await control.screenshot(path=str(OUT / "notes_menu.png"))
            await control.keyboard.press("Escape")
            await control.wait_for_selector(".menu", state="detached")
            assert not control.errors, control.errors
