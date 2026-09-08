"""
Slide Validator / QA (FIXED)

Fixed: Font size check now checks both paragraph and run level fonts.
"""

from __future__ import annotations

from typing import Dict, List, Optional, Any, Tuple
import threading

from pptx import Presentation
from pptx.util import Inches, Pt, Emu

from app.presentation.models.presentation_models import (
    SlideValidationResult,
    SlidePlan,
)
from app.presentation.design.design_system import design_system


class SlideValidator:
    """Validates rendered slides for quality issues"""
    
    def __init__(
        self,
        min_font_size: int = 10,
        max_shapes: int = 20,
        max_text_length: int = 500,
    ):
        self.min_font_size = min_font_size
        self.max_shapes = max_shapes
        self.max_text_length = max_text_length
        self._lock = threading.RLock()
    
    def validate_slide(
        self,
        slide,
        plan: Optional[SlidePlan] = None,
    ) -> SlideValidationResult:
        with self._lock:
            issues = []
            
            if len(slide.shapes) == 0:
                issues.append("Empty slide - no shapes")
                return SlideValidationResult(
                    is_valid=False, issues=issues, severity="severe"
                )
            
            if len(slide.shapes) > self.max_shapes:
                issues.append(f"Too many shapes: {len(slide.shapes)} > {self.max_shapes}")
            
            small_fonts = self._check_font_sizes(slide)
            if small_fonts:
                issues.append(f"Font too small: {len(small_fonts)} text runs below {self.min_font_size}pt")
            
            overflow = self._check_text_overflow(slide)
            if overflow:
                issues.append(f"Text overflow detected in {len(overflow)} shapes")
            
            out_of_bounds = self._check_bounds(slide)
            if out_of_bounds:
                issues.append(f"Out of bounds: {len(out_of_bounds)} shapes outside slide area")
            
            overlaps = self._check_overlap(slide)
            if overlaps:
                issues.append(f"Shape overlap detected: {len(overlaps)} pairs")
            
            if not issues:
                return SlideValidationResult(is_valid=True, issues=[], severity="none")
            
            severe_issues = [
                issue for issue in issues
                if "Empty slide" in issue
                or "Font too small" in issue
                or "Out of bounds" in issue
            ]
            
            if severe_issues:
                return SlideValidationResult(is_valid=False, issues=issues, severity="severe")
            
            return SlideValidationResult(is_valid=True, issues=issues, severity="warning")
    
    def _check_font_sizes(self, slide) -> List[str]:
        """Check for font sizes below minimum at both paragraph and run level"""
        small_fonts = []
        
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            
            for paragraph in shape.text_frame.paragraphs:
                # Check paragraph-level font
                if paragraph.font.size and paragraph.font.size < Pt(self.min_font_size):
                    small_fonts.append(paragraph.text[:30] if paragraph.text else "")
                
                # Check run-level font
                for run in paragraph.runs:
                    if run.font.size and run.font.size < Pt(self.min_font_size):
                        small_fonts.append(run.text[:30] if run.text else "")
        
        return small_fonts
    
    def _check_text_overflow(self, slide) -> List[str]:
        overflow = []
        
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            
            text = shape.text_frame.text
            if not text:
                continue
            
            shape_width = shape.width
            shape_height = shape.height
            text_length = len(text)
            
            if text_length > 200 and shape_width < Inches(2):
                overflow.append(shape.shape_id)
            elif text_length > 500 and shape_height < Inches(1):
                overflow.append(shape.shape_id)
        
        return overflow
    
    def _check_bounds(self, slide) -> List[str]:
        """Check if shapes are outside slide bounds with tolerance"""
        out_of_bounds = []
        
        slide_width = design_system.slide_width
        slide_height = design_system.slide_height
        tolerance = 0.05  # 0.05 inch tolerance
        
        for shape in slide.shapes:
            if shape.left is None or shape.top is None:
                continue
            
            left = shape.left / 914400
            top = shape.top / 914400
            width = shape.width / 914400 if shape.width else 0
            height = shape.height / 914400 if shape.height else 0
            
            right = left + width
            bottom = top + height
            
            if left < -tolerance or right > slide_width + tolerance:
                out_of_bounds.append(shape.shape_id)
            elif top < -tolerance or bottom > slide_height + tolerance:
                out_of_bounds.append(shape.shape_id)
        
        return out_of_bounds
    
    def _check_overlap(self, slide) -> List[Tuple[int, int]]:
        overlaps = []
        shapes = list(slide.shapes)
        
        for i in range(len(shapes)):
            for j in range(i + 1, len(shapes)):
                shape_a = shapes[i]
                shape_b = shapes[j]
                
                if self._shapes_overlap(shape_a, shape_b):
                    overlaps.append((shape_a.shape_id, shape_b.shape_id))
        
        return overlaps
    
    def _shapes_overlap(self, shape_a, shape_b) -> bool:
        if shape_a.left is None or shape_b.left is None:
            return False
        if shape_a.top is None or shape_b.top is None:
            return False
        
        a_left = shape_a.left / 914400
        a_top = shape_a.top / 914400
        a_right = a_left + (shape_a.width / 914400 if shape_a.width else 0)
        a_bottom = a_top + (shape_a.height / 914400 if shape_a.height else 0)
        
        b_left = shape_b.left / 914400
        b_top = shape_b.top / 914400
        b_right = b_left + (shape_b.width / 914400 if shape_b.width else 0)
        b_bottom = b_top + (shape_b.height / 914400 if shape_b.height else 0)
        
        return not (
            a_right <= b_left
            or a_left >= b_right
            or a_bottom <= b_top
            or a_top >= b_bottom
        )