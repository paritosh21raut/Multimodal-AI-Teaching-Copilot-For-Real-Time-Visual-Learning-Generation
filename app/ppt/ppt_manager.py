from __future__ import annotations

from datetime import datetime
from pathlib import Path
from threading import Lock
from typing import Any, Dict, List, Optional

from pptx import Presentation
from pptx.slide import Slide

from app.slides.slide_renderer import slide_renderer


class PPTManager:
    """
    Manages one PowerPoint presentation for a lecture.
    """

    def __init__(self) -> None:

        self._lock = Lock()

        self.presentation: Optional[
            Presentation
        ] = None

        self.presentation_path: Optional[
            Path
        ] = None

        self.output_directory = Path(
            "outputs/presentations"
        )

        self.output_directory.mkdir(
            parents=True,
            exist_ok=True,
        )

        self.slide_map: Dict[
            int,
            Slide,
        ] = {}

    # ============================================================
    # PRESENTATION
    # ============================================================

    def create_new_presentation(
        self,
        lecture_title: str = "Lecture",
    ) -> Path:

        with self._lock:

            self.presentation = (
                Presentation()
            )

            self.slide_map.clear()

            timestamp = datetime.now().strftime(
                "%Y%m%d_%H%M%S"
            )

            filename = (
                lecture_title.replace(
                    " ",
                    "_",
                )
                + "_"
                + timestamp
                + ".pptx"
            )

            self.presentation_path = (
                self.output_directory
                / filename
            )

            self.save()

            print(
                f"[PPT] Created: "
                f"{self.presentation_path}"
            )

            return self.presentation_path

    # ============================================================
    # TITLE SLIDE
    # ============================================================

    def add_title_slide(
        self,
        title: str,
        subtitle: str = "",
    ) -> None:

        with self._lock:

            self._ensure_presentation()

            layout = (
                self.presentation
                .slide_layouts[0]
            )

            slide = (
                self.presentation
                .slides
                .add_slide(layout)
            )

            slide_renderer.render_title_slide(
                slide=slide,
                title=title,
                subtitle=subtitle,
            )

            self.save()

            print(
                "[PPT] Title slide created"
            )

    # ============================================================
    # CONTENT SLIDE
    # ============================================================

    def create_or_update_slide(
        self,
        slide_id: int,
        title: str,
        bullets: List[str],
        image_path: str | None = None,
        visual_type: str = "none",
        visual_spec: Optional[
            Dict[str, Any]
        ] = None,
        content_type: str = "explanation",
        visual_reason: str = "",
    ):

        print(
            "[PPT] presentation =",
            self.presentation,
        )

        print(
            "[PPT] path =",
            self.presentation_path,
        )

        with self._lock:

            self._ensure_presentation()

            if visual_spec is None:
                visual_spec = {}

            if slide_id in self.slide_map:

                slide = self.slide_map[
                    slide_id
                ]

                self._update_slide(
                    slide=slide,
                    title=title,
                    bullets=bullets,
                    image_path=image_path,
                    visual_type=visual_type,
                    visual_spec=visual_spec,
                    content_type=content_type,
                    visual_reason=visual_reason,
                )

                print(
                    f"[PPT] Updated Slide "
                    f"{slide_id}"
                )

            else:

                layout = (
                    self.presentation
                    .slide_layouts[1]
                )

                slide = (
                    self.presentation
                    .slides
                    .add_slide(layout)
                )

                self.slide_map[
                    slide_id
                ] = slide

                self._update_slide(
                    slide=slide,
                    title=title,
                    bullets=bullets,
                    image_path=image_path,
                    visual_type=visual_type,
                    visual_spec=visual_spec,
                    content_type=content_type,
                    visual_reason=visual_reason,
                )

                print(
                    f"[PPT] Created Slide "
                    f"{slide_id}"
                )

            self.save()

    # ============================================================
    # INTERNAL
    # ============================================================

    def _update_slide(
        self,
        slide,
        title,
        bullets,
        image_path=None,
        visual_type="none",
        visual_spec=None,
        content_type="explanation",
        visual_reason="",
    ):

        slide_renderer.render_content_slide(
            slide=slide,
            title=title,
            bullets=bullets,
            image_path=image_path,
            visual_type=visual_type,
            visual_spec=visual_spec or {},
            content_type=content_type,
            visual_reason=visual_reason,
        )

    # ============================================================
    # SAVE
    # ============================================================

    def save(self) -> None:

        self._ensure_presentation()

        self.presentation.save(
            self.presentation_path
        )

    # ============================================================
    # GETTERS
    # ============================================================

    def get_presentation(
        self,
    ) -> Presentation:

        self._ensure_presentation()

        return self.presentation

    def get_path(
        self,
    ) -> Optional[Path]:

        return self.presentation_path

    def slide_exists(
        self,
        slide_id: int,
    ) -> bool:

        return (
            slide_id
            in self.slide_map
        )

    def total_slides(
        self,
    ) -> int:

        return len(
            self.slide_map
        )

    # ============================================================
    # VALIDATION
    # ============================================================

    def _ensure_presentation(
        self,
    ) -> None:

        if self.presentation is None:

            raise RuntimeError(
                "Presentation has not been created."
            )


ppt_manager = PPTManager()