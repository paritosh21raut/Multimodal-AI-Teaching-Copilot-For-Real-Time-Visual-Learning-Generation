"""PDF and PPTX files for the lecture materials (F-010), rendered by headless Edge through our own server.

- PDF (notes, assignment): print HTML built here; formulas with the vendored KaTeX; Edge prints it (A4).
- PPTX: every slide drawn by the real slide renderer (`/web/export/`) at 1920×1080 in the chosen theme. Each plain text
  element becomes an editable text box at its rendered place, font, size and colour; the rest (cards, arrows, images,
  formulas) is the slide's background picture, drawn with those texts hidden. 1 stage pixel = 0.5 pt.
"""
from __future__ import annotations

import html
import logging
import re
from pathlib import Path
from typing import Any, Optional, Sequence

from copilot.materials.content import Lecture, Subtopic
from copilot.materials.writer import Concept, Question
from copilot.presentation.mathtext import chem_parts

log = logging.getLogger(__name__)

_CHEM_TOKEN = re.compile(r"(?<![A-Za-z0-9])(\d*(?:[A-Z][a-z]?\d*)+)(?![A-Za-z0-9])")
_MATH_SPAN = re.compile(r"\\\((.+?)\\\)")


def rich(text: str) -> str:
    """Escaped text with chemical formulas subscripted (CO2 → CO<sub>2</sub>), as on the slides."""
    out, last = [], 0
    for m in _CHEM_TOKEN.finditer(text or ""):
        tok = m.group(1)
        coef = re.match(r"\d*", tok).group(0)  # type: ignore[union-attr]
        parts = chem_parts(tok)
        if not parts:
            continue
        out.append(html.escape(text[last:m.start()]) + html.escape(coef))
        out.append("".join(html.escape(s) + (f"<sub>{n}</sub>" if n else "") for s, n in parts))
        last = m.end()
    out.append(html.escape((text or "")[last:]))
    return "".join(out)


def mixed(text: str, math: str = "") -> str:
    """Slide text whose equations are marked \\( latex \\): each equation rendered by KaTeX in the page."""
    if not math:
        return rich(text)
    out, last = [], 0
    for m in _MATH_SPAN.finditer(math):
        out.append(rich(math[last:m.start()]))
        out.append(f'<span class="tex" data-tex="{html.escape(m.group(1))}">{html.escape(m.group(1))}</span>')
        last = m.end()
    out.append(rich(math[last:]))
    return "".join(out)


def _tree(node: dict[str, Any]) -> str:
    kids = node.get("children") or []
    inner = "".join(_tree(c) for c in kids)
    return f"<li>{rich(node.get('label', ''))}{f'<ul>{inner}</ul>' if inner else ''}</li>"


def block_html(b: dict[str, Any]) -> str:
    """A slide block as print HTML (images are handled per subtopic)."""
    t = b.get("type")
    if t == "definition":
        notes = "".join(f"<li>{mixed(n['text'], n.get('math', ''))}</li>" for n in b.get("notes") or [])
        return (f'<div class="def"><span class="def-term">{rich(b["term"])}</span>'
                f'<p>{mixed(b["definition"], b.get("math", ""))}</p>{f"<ul>{notes}</ul>" if notes else ""}</div>')
    if t == "points":
        tag = "ol" if b.get("style") == "numbers" else "ul"
        head = f'<p class="lead">{rich(b["heading"])}</p>' if b.get("heading") else ""
        items = "".join(f"<li>{mixed(i['text'], i.get('math', ''))}</li>" for i in b.get("items") or [])
        return f"{head}<{tag}>{items}</{tag}>"
    if t == "process":
        steps = "".join(f"<li><b>{mixed(s['label'], s.get('math', ''))}</b>"
                        f"{(' — ' + rich(s['detail'])) if s.get('detail') else ''}</li>" for s in b.get("steps") or [])
        cyc = '<p class="small">The steps repeat as a cycle.</p>' if b.get("cyclic") else ""
        return f'<ol class="steps" start="{b.get("start") or 1}">{steps}</ol>{cyc}'
    if t == "comparison":
        head = "".join(f"<th>{rich(c['heading'])}</th>" for c in b.get("columns") or [])
        rows = "".join(f"<tr><th>{rich(r['aspect'])}</th>" + "".join(f"<td>{rich(c)}</td>" for c in r["cells"])
                       + "</tr>" for r in b.get("rows") or [])
        return f'<table class="compare"><thead><tr><th></th>{head}</tr></thead><tbody>{rows}</tbody></table>'
    if t == "timeline":
        ev = "".join(f"<li><b>{rich(e['when'])}</b> {rich(e['label'])}"
                     f"{(' — ' + rich(e['detail'])) if e.get('detail') else ''}</li>" for e in b.get("events") or [])
        return f'<ul class="timeline">{ev}</ul>'
    if t == "hierarchy":
        return f'<ul class="tree">{_tree(b["root"])}</ul>'
    if t == "cause_effect":
        links = "".join(f"<li>{rich(l['cause'])} <span class='arrow'>→</span> {rich(l['effect'])}</li>"
                        for l in b.get("links") or [])
        return f'<ul class="causes">{links}</ul>'
    if t == "formula":
        eq = (f'<div class="formula"><span class="tex display" data-tex="{html.escape(b["latex"])}">'
              f'{html.escape(b.get("spoken") or b["latex"])}</span></div>' if b.get("latex")
              else f'<div class="formula">{rich(b.get("spoken", ""))}</div>')
        def symbol(v: dict[str, Any]) -> str:
            if not v.get("latex"):
                return rich(v["symbol"])
            return f'<span class="tex" data-tex="{html.escape(v["latex"])}">{html.escape(v["symbol"])}</span>'
        variables = b.get("variables") or []
        units = any(v.get("unit") for v in variables)  # no empty unit column

        def row(v: dict[str, Any]) -> str:
            unit = f"<td>{rich(v.get('unit', ''))}</td>" if units else ""
            return f"<tr><td>{symbol(v)}</td><td>{rich(v['meaning'])}</td>{unit}</tr>"
        rows = "".join(row(v) for v in variables)
        return eq + (f'<table class="vars"><tbody>{rows}</tbody></table>' if rows else "")
    if t == "example":
        title = f"<b>{rich(b['title'])}</b> " if b.get("title") else ""
        return f'<div class="example"><span class="tag">Example</span><p>{title}{mixed(b["text"], b.get("math", ""))}</p></div>'
    if t == "callout":
        return f'<div class="callout"><p>{mixed(b["text"], b.get("math", ""))}</p></div>'
    if t == "facts":
        head = f'<p class="lead">{rich(b["heading"])}</p>' if b.get("heading") else ""
        rows = "".join(f"<tr><th>{rich(f['label'])}</th><td>{rich(f.get('value', ''))}</td></tr>"
                       for f in b.get("facts") or [])
        return f'{head}<table class="facts"><tbody>{rows}</tbody></table>'
    if t == "groups":
        head = f'<p class="lead">{rich(b["heading"])}</p>' if b.get("heading") else ""
        groups = "".join(f"<li><b>{rich(g['label'])}:</b> " + "; ".join(rich(i["text"]) for i in g.get("items") or [])
                         + "</li>" for g in b.get("groups") or [])
        return f"{head}<ul>{groups}</ul>"
    return ""


def _figure(sub: Subtopic) -> str:
    for s in sub.slides:
        for b in s.get("blocks") or []:
            if b.get("type") == "image" and b.get("url", "").startswith("/media/"):
                credit = " · ".join(x for x in (b.get("credit", ""), b.get("licence", "")) if x)
                return (f'<figure><img src="{html.escape(b["url"])}" alt="{html.escape(b.get("alt", ""))}">'
                        f'<figcaption>{rich(b.get("alt", "").capitalize())}'
                        f'{f"<span>{html.escape(credit[:160])}</span>" if credit else ""}</figcaption></figure>')
    return ""


def _sub_html(sub: Subtopic, text: str) -> str:
    seen: set[str] = set()
    blocks = []
    for s in sub.slides:
        for b in s.get("blocks") or []:
            h = block_html(b)
            if h and h not in seen:
                seen.add(h)
                blocks.append(h)
    explain = f'<p class="explain">{rich(text)}</p>' if text else ""
    heading = "" if sub.title.lower() == sub.topic.lower() else f"<h3>{rich(sub.title)}</h3>"
    return f'<article class="sub">{heading}{explain}{_figure(sub)}{"".join(blocks)}</article>'


def notes_html(title: str, lectures: Sequence[Lecture], sections: dict[str, str],
               concepts: Sequence[Concept]) -> str:
    body = []
    many = len(lectures) > 1
    toc = []
    for n, lec in enumerate(lectures, start=1):
        if many:
            body.append(f'<div class="part"><span>Lecture {n}</span>{rich(lec.title)}'
                        f'{f"<em>{lec.date}</em>" if lec.date else ""}</div>')
        for t in lec.topics:
            toc.append(f"<li>{rich(t.name)}</li>")
            subs = "".join(_sub_html(s, sections.get(s.key, "")) for s in t.subtopics)
            body.append(f'<section class="topic"><h2>{rich(t.name)}</h2>{subs}</section>')
    if concepts:
        rows = "".join(f"<div class='term'><dt>{rich(c.term)}</dt><dd>{rich(c.meaning)}</dd></div>" for c in concepts)
        body.append(f'<section class="glossary"><h2>Key concepts</h2><dl>{rows}</dl></section>')
    dates = ", ".join(l.date for l in lectures if l.date)
    contents = f'<ol class="toc">{"".join(toc)}</ol>' if len(toc) > 2 else ""
    head = (f'<header class="doc-head"><div class="kicker">Study notes</div><h1>{rich(title)}</h1>'
            f'<div class="meta">{html.escape(dates)}</div>{contents}</header>')
    return _page(title, head + "".join(body))


SHORT_MARKS, LONG_MARKS = 2, 5


def assignment_html(title: str, lectures: Sequence[Lecture], questions: Sequence[Question]) -> str:
    short = [q for q in questions if q.type == "short"]
    long_ = [q for q in questions if q.type == "long"]
    total = SHORT_MARKS * len(short) + LONG_MARKS * len(long_)
    dates = ", ".join(l.date for l in lectures if l.date)
    parts, n = [], 1
    for label, qs, marks, hint in (("A", short, SHORT_MARKS, "Short answers"), ("B", long_, LONG_MARKS, "Long answers")):
        if not qs:
            continue
        items = []
        for q in qs:
            items.append(f'<li value="{n}"><p>{rich(q.text)}</p><span class="marks">{marks}</span></li>')
            n += 1
        parts.append(f'<section class="qs"><h2>Section {label} <span>{hint} · {marks} marks each</span></h2>'
                     f'<ol>{"".join(items)}</ol></section>')
    head = (f'<header class="doc-head"><div class="kicker">Assignment</div><h1>{rich(title)}</h1>'
            f'<div class="meta">{len(questions)} questions · {total} marks{f" · {html.escape(dates)}" if dates else ""}'
            f'</div></header><div class="student"><span>Name</span><span>Roll no.</span><span>Date</span></div>'
            f'<p class="instructions">Answer all questions. Write in your own words.</p>')
    return _page(title, head + "".join(parts))


PRINT_CSS = """
@page { size: A4; margin: 18mm 17mm 20mm; }
* { box-sizing: border-box; }
html { -webkit-print-color-adjust: exact; print-color-adjust: exact; }
body { margin: 0; font: 10.5pt/1.55 "Segoe UI Variable Text", "Segoe UI", system-ui, sans-serif; color: #16191d; }
h1, h2, h3, .kicker, .part, dt { font-family: "Segoe UI Variable Display", "Segoe UI", system-ui, sans-serif; }
.doc-head { padding: 4mm 0 6mm; border-bottom: 2px solid #0e6e66; margin-bottom: 7mm; }
.kicker { color: #0e6e66; font-weight: 700; letter-spacing: .14em; text-transform: uppercase; font-size: 9pt; }
h1 { font-size: 26pt; line-height: 1.12; margin: 2mm 0 1mm; letter-spacing: -0.01em; }
.meta { color: #6b7380; font-size: 9.5pt; }
.toc { columns: 2; margin: 5mm 0 0; padding-left: 5mm; color: #3d444d; font-size: 9.5pt; }
.part { margin: 9mm 0 4mm; padding: 3mm 4mm; background: #0e6e66; color: #fff; border-radius: 6px; font-size: 14pt;
  font-weight: 650; break-after: avoid; }
.part span { display: block; font-size: 8pt; letter-spacing: .14em; text-transform: uppercase; opacity: .8; }
.part em { font-style: normal; font-size: 9pt; opacity: .8; margin-left: 3mm; }
h2 { font-size: 15.5pt; color: #0e6e66; margin: 8mm 0 3mm; break-after: avoid; }
h3 { font-size: 12pt; margin: 5mm 0 1.5mm; break-after: avoid; }
.topic > .sub:first-of-type > h3 { margin-top: 2mm; }
p { margin: 0 0 2.5mm; }
.explain { color: #22272e; }
ul, ol { margin: 0 0 3mm; padding-left: 6mm; }
li { margin: 0 0 1mm; }
.lead { font-weight: 650; margin-top: 2mm; }
.def { border-left: 3px solid #0e6e66; background: #eef6f4; padding: 2.5mm 4mm; border-radius: 0 6px 6px 0;
  margin: 2mm 0 3mm; break-inside: avoid; }
.def-term { font-weight: 700; color: #0e6e66; }
.def p { margin: .5mm 0 0; }
.steps li::marker { color: #0e6e66; font-weight: 700; }
table { border-collapse: collapse; width: 100%; margin: 2mm 0 4mm; font-size: 9.5pt; break-inside: avoid; }
th, td { border: 1px solid #ddd6c8; padding: 1.6mm 2.5mm; text-align: left; vertical-align: top; }
thead th { background: #efebe2; }
tbody th { background: #f8f6f1; font-weight: 600; }
.vars { width: auto; min-width: 60%; }
.formula { text-align: center; font-size: 13pt; margin: 3mm 0; padding: 3mm; background: #f8f6f1; border-radius: 6px;
  break-inside: avoid; }
.example, .callout { background: #f6e3d6; border-radius: 6px; padding: 2.5mm 4mm; margin: 2mm 0 3mm;
  break-inside: avoid; }
.callout { background: #efebe2; }
.tag { font-size: 7.5pt; font-weight: 700; letter-spacing: .12em; text-transform: uppercase; color: #b4531f; }
.example p, .callout p { margin: .5mm 0 0; }
.arrow { color: #0e6e66; font-weight: 700; }
.tree ul { margin-bottom: 0; }
figure { margin: 2mm 0 4mm; text-align: center; break-inside: avoid; }
figure img { max-width: 100%; max-height: 62mm; border-radius: 6px; }
figcaption { font-size: 8.5pt; color: #6b7380; margin-top: 1mm; }
figcaption span { display: block; font-size: 7pt; }
.glossary { break-before: auto; }
.glossary dl { margin: 0; display: grid; grid-template-columns: 1fr 1fr; gap: 2.5mm 6mm; }
.term { break-inside: avoid; border-top: 2px solid #0e6e66; padding-top: 1.5mm; }
dt { font-weight: 700; color: #0e6e66; }
dd { margin: 0; }
.small { font-size: 9pt; color: #6b7380; }
.student { display: grid; grid-template-columns: 2fr 1fr 1fr; gap: 6mm; margin: 0 0 5mm; }
.student span { border-bottom: 1px solid #9aa1ab; padding: 5mm 0 1mm; color: #6b7380; font-size: 9pt; }
.instructions { color: #3d444d; font-style: italic; }
.qs h2 span { font-size: 9.5pt; color: #6b7380; font-weight: 500; margin-left: 2mm; }
.qs ol { padding-left: 8mm; }
.qs li { position: relative; padding-right: 14mm; margin: 0 0 4mm; break-inside: avoid; }
.qs li::marker { font-weight: 700; color: #0e6e66; }
.qs li p { margin: 0; }
.marks { position: absolute; right: 0; top: 0; font-size: 8.5pt; color: #6b7380; border: 1px solid #ddd6c8;
  border-radius: 99px; padding: 0 2mm; }
.marks::after { content: " m"; }
sub { font-size: 70%; }
"""

KATEX_SCRIPT = """<script type="module">
import katex from "/web/vendor/katex/katex.mjs";
for (const el of document.querySelectorAll("[data-tex]")) {
  try { el.innerHTML = katex.renderToString(el.dataset.tex, { throwOnError: true, displayMode: el.classList.contains("display"), strict: "ignore" }); }
  catch (e) { /* not renderable: the plain text stays */ }
}
await document.fonts.ready;
await Promise.all([...document.images].map((i) => i.complete ? 0 : new Promise((r) => { i.onload = i.onerror = r; })));
window.__ready = true;
</script>"""


def _page(title: str, body: str) -> str:
    return (f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{html.escape(title)}</title>'
            f'<link rel="stylesheet" href="/web/vendor/katex/katex.min.css"><style>{PRINT_CSS}</style></head>'
            f"<body>{body}{KATEX_SCRIPT}</body></html>")


FOOTER = ('<div style="width:100%;font:8px Segoe UI,sans-serif;color:#9aa1ab;padding:0 17mm;display:flex;'
          'justify-content:space-between"><span>{title}</span><span><span class="pageNumber"></span> / '
          '<span class="totalPages"></span></span></div>')

# ---- PPTX -----------------------------------------------------------------------------------------------------------
STAGE_W, STAGE_H = 1920, 1080
EMU_PER_PX = 12192000 / STAGE_W  # 13.333 in wide: 1 stage px = 6350 EMU = 0.5 pt

# In the export page: the stage's text grouped by the box it flows in (its nearest non-inline ancestor: a card, a
# badge, a heading — never a badge merged into the text beside it), with its runs and styles and the tight union of
# its line boxes. A box whose flow holds a formula stays in the background. Then those texts are hidden so the
# screenshot is the background without them.
COLLECT_JS = """() => {
  const stage = document.querySelector(".stage");
  const sr = stage.getBoundingClientRect();
  const blockOf = (el) => { while (el !== stage && getComputedStyle(el).display === "inline") el = el.parentElement; return el; };
  const groups = new Map();
  const skip = new Set();
  for (const t of stage.querySelectorAll(".katex, .tex, .tex-fallback")) skip.add(blockOf(t.parentElement));
  const walker = document.createTreeWalker(stage, NodeFilter.SHOW_TEXT);
  for (let n = walker.nextNode(); n; n = walker.nextNode()) {
    if (!n.data.trim() || n.parentElement.closest(".katex, .tex, svg")) continue;
    const block = blockOf(n.parentElement);
    if (skip.has(block)) continue;
    const cs = getComputedStyle(n.parentElement);
    if (cs.visibility === "hidden" || +cs.opacity === 0) continue;
    if (!groups.has(block)) groups.set(block, []);
    groups.get(block).push(n);
  }
  const boxes = [];
  for (const [block, nodes] of groups) {
    const rects = [];
    const runs = [];
    for (const n of nodes) {
      const range = document.createRange(); range.selectNodeContents(n);
      rects.push(...[...range.getClientRects()].filter((q) => q.width > 0 && q.height > 0));
      const s = getComputedStyle(n.parentElement);
      let text = n.data.replace(/\\s+/g, " ");
      if (s.textTransform === "uppercase") text = text.toUpperCase();
      runs.push({ text, size: parseFloat(s.fontSize), weight: +s.fontWeight || 400, italic: s.fontStyle === "italic",
        color: s.color, spacing: parseFloat(s.letterSpacing) || 0, sub: n.parentElement.tagName === "SUB",
        sup: n.parentElement.tagName === "SUP" });
    }
    if (!rects.length) continue;
    const x0 = Math.min(...rects.map((q) => q.left)), y0 = Math.min(...rects.map((q) => q.top));
    const x1 = Math.max(...rects.map((q) => q.right)), y1 = Math.max(...rects.map((q) => q.bottom));
    const cs = getComputedStyle(block);
    const lines = new Set(rects.map((q) => Math.round(q.top / 4))).size;
    boxes.push({ x: x0 - sr.left, y: y0 - sr.top, w: x1 - x0, h: y1 - y0, lines, align: cs.textAlign,
      lineHeight: parseFloat(cs.lineHeight) || parseFloat(cs.fontSize) * 1.2, family: cs.fontFamily, runs });
    for (const n of nodes) {  // hide exactly these texts: an inline wrapper changes no layout, hides no other child
      const hide = document.createElement("span");
      hide.style.cssText = "color:transparent!important;-webkit-text-fill-color:transparent!important;text-shadow:none!important";
      n.parentNode.insertBefore(hide, n);
      hide.appendChild(n);
    }
  }
  return boxes;
}"""


def _rgb(css: str) -> tuple[int, int, int, float]:
    nums = [float(x) for x in re.findall(r"[\d.]+", css)]
    if len(nums) < 3:
        return 0, 0, 0, 1.0
    return int(nums[0]), int(nums[1]), int(nums[2]), nums[3] if len(nums) > 3 else 1.0


def _font(weight: int, family: str) -> tuple[str, bool]:
    """Segoe UI faces PowerPoint has on every Windows machine for the slide's Segoe UI Variable weights."""
    if "mono" in family.lower() or "consolas" in family.lower():
        return "Consolas", weight >= 600
    if weight >= 620:  # the slides' 650 display weight looks bold (PowerPoint export compared 2026-10-07)
        return "Segoe UI", True
    if weight >= 520:
        return "Segoe UI Semibold", False
    if weight <= 350:
        return "Segoe UI Light", False
    return "Segoe UI", False


def add_text_box(slide, box: dict[str, Any], bg: tuple[int, int, int]) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Pt

    slack = max(6.0, box["w"] * 0.06) if box["lines"] > 1 or box["w"] > 200 else 6.0
    x, w = box["x"], box["w"] + slack
    if box["align"] == "center":
        x -= slack / 2
    elif box["align"] in ("right", "end"):
        x -= slack
    tb = slide.shapes.add_textbox(Emu(int(x * EMU_PER_PX)), Emu(int(box["y"] * EMU_PER_PX)),
                                  Emu(int(w * EMU_PER_PX)), Emu(int(box["h"] * EMU_PER_PX)))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT}.get(box["align"],
                                                                                                  PP_ALIGN.LEFT)
    if box["lines"] > 1:  # wrapped text: the slide's own line spacing keeps every line where the browser put it
        tf.vertical_anchor = MSO_ANCHOR.TOP
        p.line_spacing = Pt(box["lineHeight"] * 0.5)
    else:  # one line: centred on its line box (an exact tight spacing pushed a badge's "I" to the top)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    runs = box["runs"]
    if runs:  # the container's own leading / trailing spaces are layout, not text
        runs[0] = {**runs[0], "text": runs[0]["text"].lstrip()}
        runs[-1] = {**runs[-1], "text": runs[-1]["text"].rstrip()}
    for r in runs:
        if not r["text"]:
            continue
        run = p.add_run()
        run.text = r["text"]
        name, bold = _font(r["weight"], box["family"])
        f = run.font
        f.name, f.bold, f.italic = name, bold, r["italic"]
        f.size = Pt(round(r["size"] * 0.5 * (1 / 0.7 if r["sub"] or r["sup"] else 1), 1))
        red, green, blue, alpha = _rgb(r["color"])
        if alpha < 1:  # blend a translucent colour onto the slide background
            red, green, blue = (round(c * alpha + b * (1 - alpha)) for c, b in zip((red, green, blue), bg))
        f.color.rgb = RGBColor(red, green, blue)
        rpr = run._r.get_or_add_rPr()
        if r["spacing"]:
            rpr.set("spc", str(int(r["spacing"] * 0.5 * 100)))
        if r["sub"] or r["sup"]:
            rpr.set("baseline", "-25000" if r["sub"] else "30000")


class Renderer:
    """Headless Edge on our own server (`base_url`, a local address: the teacher's pages and /media)."""

    def __init__(self, base_url: str) -> None:
        self.base_url = base_url.rstrip("/")

    async def _browser(self, pw):
        return await pw.chromium.launch(channel="msedge", headless=True)

    async def pdf(self, page_html: str, out: Path, footer_title: str) -> None:
        from playwright.async_api import async_playwright

        async with async_playwright() as pw:
            browser = await self._browser(pw)
            try:
                page = await browser.new_page()
                await page.goto(f"{self.base_url}/web/export/blank.html")
                await page.set_content(page_html, wait_until="load")
                await page.wait_for_function("window.__ready === true", timeout=20000)
                await page.pdf(path=str(out), format="A4", print_background=True, prefer_css_page_size=True,
                               display_header_footer=True, header_template="<span></span>",
                               footer_template=FOOTER.replace("{title}", html.escape(footer_title[:90])))
            finally:
                await browser.close()

    async def pptx(self, slides: Sequence[dict[str, Any]], theme: str, out: Path, title: str = "",
                   frames_dir: Optional[Path] = None) -> int:
        """The slides as a .pptx in `theme`; frames_dir: also keep each slide's PNG (tests, checking)."""
        import io

        from playwright.async_api import async_playwright
        from pptx import Presentation
        from pptx.util import Emu

        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
        blank = prs.slide_layouts[6]
        if title:
            prs.core_properties.title = title
        async with async_playwright() as pw:
            browser = await self._browser(pw)
            try:
                page = await browser.new_page(viewport={"width": STAGE_W, "height": STAGE_H})
                await page.goto(f"{self.base_url}/web/export/index.html")
                await page.wait_for_function("window.__ready === true", timeout=15000)
                for n, spec in enumerate(slides):
                    await page.evaluate("([spec, theme]) => window.showSlide(spec, theme)", [spec, theme])
                    if frames_dir is not None:
                        await page.screenshot(path=str(frames_dir / f"slide{n + 1:02d}_full.png"))
                    boxes = await page.evaluate(COLLECT_JS)
                    bg_css = await page.evaluate(
                        "getComputedStyle(document.querySelector('.stage')).backgroundColor")
                    png = await page.screenshot(type="png")
                    if frames_dir is not None:
                        (frames_dir / f"slide{n + 1:02d}_bg.png").write_bytes(png)
                    s = prs.slides.add_slide(blank)
                    s.shapes.add_picture(io.BytesIO(png), 0, 0, prs.slide_width, prs.slide_height)
                    bg = _rgb(bg_css)[:3]
                    for box in boxes:
                        add_text_box(s, box, bg)
                    if spec.get("title"):
                        s.notes_slide.notes_text_frame.text = spec["title"]
            finally:
                await browser.close()
        prs.save(str(out))
        return len(slides)
